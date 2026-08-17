"""旧格式工作簿适配器（risks.md R8 建议 1、data-survey.md §4）。

把 `data/` 下 11 表、4 行多级表头的处理后工作簿转换为 `TfdcDataset`
内存对象（不落 xlsx），产出走 `importer.import_parsed` 标准管线（含全部校验）。

表头结构（data-survey §2）::

    r1 设备类        冷水主机
    r2 设备实例      chiller_01        （部分表标作「设备名称」，不依赖标签）
    r3 物模型        chiller
    r4 物模型属性    power | 电流百分比 | ...
    r5+ Excel 序列号时间戳（naive） | 数据 ...

- timestamp 为 Excel 序列号（naive），按 manifest.timezone=Asia/Shanghai
  显式声明解释，降级行为记入 `LegacyConversion.pre_diagnostics`（TFDC-502 WARN）。
- 中文属性名映射为 lower_snake_case（`电流百分比` → `current_percent`），
  映射表记录进 lineage。
- 物模型名 → 注册 TFOM 的映射见 `MODEL_ID_MAP`。
- 布尔属性（TFOM dtype=boolean）的 0/1 单元格在适配层转为 False/True；
  其他取值（如 19）原样传入管线，以 TFDC-404 REJECT 呈现（data-survey §F7）。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from thermoforge_core.contracts.tfdc import (
    ObjectRecord,
    TfdcDataset,
    TfdcManifest,
    VariableRecord,
)
from thermoforge_core.errors import Diagnostic, Level
from thermoforge_core.naming import is_property_code

from .importer import TfomRegistry

LEGACY_TIMEZONE = "Asia/Shanghai"

# 中文/非常规属性名 → lower_snake_case property_code（记录进 lineage）
PROPERTY_NAME_MAP: dict[str, str] = {
    "电流百分比": "current_percent",
}

# r3 物模型名 → 注册的 object_model_id
MODEL_ID_MAP: dict[str, str] = {
    "chiller": "chiller.v2",
    "chilled_water_pump": "chilled_water_pump.v1",
    "cooling_water_pump": "cooling_water_pump.v1",
    "cooling_tower_fan": "cooling_tower_fan.v1",
    "cooling_tower": "cooling_tower.v1",
    "environment_parameters": "environment_parameters.v1",
    "chilled_water_header": "chilled_water_header.v1",
    "cooling_water_header": "cooling_water_header.v1",
    "heat_exchanger": "heat_exchanger.v1",
    "surface_type_heat_exchanger": "surface_type_heat_exchanger.v1",
    "lumped_load": "lumped_load.v1",
}


@dataclass
class LegacyConversion:
    """旧格式转换结果（内存对象，不落 xlsx）。"""

    dataset: TfdcDataset
    timestamps: list[datetime]  # UTC，升序，全表并集时间轴
    columns: dict[str, list[Any]]  # variable_id → 对齐到并集时间轴的值
    property_mapping: dict[str, str]  # 原始属性名 → property_code（仅映射过的）
    pre_diagnostics: list[Diagnostic] = field(default_factory=list)
    degradations: list[str] = field(default_factory=list)

    def lineage(self, source_path: str | Path) -> dict[str, Any]:
        """写入 vault lineage.json 的适配层信息。"""
        return {
            "adapter": "thermoforge_data.legacy",
            "adapter_note": "旧格式 11 表多级表头工作簿 → TfdcDataset 内存对象",
            "property_name_mapping": dict(self.property_mapping),
            "model_id_mapping": dict(MODEL_ID_MAP),
            "timezone_declaration": LEGACY_TIMEZONE,
        }


def convert_legacy_workbook(
    path: str | Path,
    registry: TfomRegistry,
    *,
    dataset_id: str = "WX_2025_HVAC",
    site_id: str = "WX",
    time_resolution: str = "900s",
) -> LegacyConversion:
    """流式读取旧格式工作簿并转换（read_only + data_only）。

    原始工作簿只读；本函数不写入任何文件。
    """
    path = Path(path)
    sheets = load_workbook_model(path)
    return convert_legacy_tables(
        sheets, registry,
        dataset_id=dataset_id, site_id=site_id,
        time_resolution=time_resolution, source_name=path.name,
    )


@dataclass
class LegacySheet:
    """旧格式工作簿的一张表（内存表示，预处理执行器的操作对象）。"""

    name: str
    header_rows: list[list[Any]]  # 4 行多级表头
    data_rows: list[tuple[Any, ...]]  # 第 5 行起


def load_workbook_model(path: str | Path) -> dict[str, LegacySheet]:
    """把旧格式工作簿流式读入内存表模型（原始文件只读，不写盘）。"""
    wb = load_workbook(Path(path), read_only=True, data_only=True,
                       keep_links=False)
    try:
        sheets: dict[str, LegacySheet] = {}
        for name in wb.sheetnames:
            rows_iter = wb[name].iter_rows(values_only=True)
            header = [list(r) for _, r in zip(range(4), rows_iter)]
            data_rows = [tuple(r) for r in rows_iter]
            sheets[name] = LegacySheet(name=name, header_rows=header,
                                       data_rows=data_rows)
        return sheets
    finally:
        wb.close()


def convert_legacy_tables(
    sheets: dict[str, LegacySheet],
    registry: TfomRegistry,
    *,
    dataset_id: str = "WX_2025_HVAC",
    site_id: str = "WX",
    time_resolution: str = "900s",
    source_name: str | None = None,
) -> LegacyConversion:
    """内存表模型 → TfdcDataset（预处理执行器与 xlsx 入口共用）。"""
    tz = ZoneInfo(LEGACY_TIMEZONE)

    # 各表列结构（r1–r4 多级表头）
    sheet_cols: dict[str, list[tuple[int, str, str, str]]] = {}
    objects: dict[str, tuple[str, str]] = {}  # instance → (model_id, class_zh)
    property_mapping: dict[str, str] = {}  # 原始属性名 → property_code
    for sheet, model in sheets.items():
        header_rows = model.header_rows
        cols, mapping = _parse_multi_header(header_rows)
        sheet_cols[sheet] = cols
        property_mapping.update(mapping)
        for _, instance, model_name, _prop in cols:
            model_id = MODEL_ID_MAP.get(model_name)
            if model_id is None:
                raise ValueError(
                    f"物模型名未在 MODEL_ID_MAP 登记: {model_name!r}（{sheet}）"
                )
            class_zh = next(
                (str(v) for v in header_rows[0][1:] if v is not None), sheet
            )
            objects.setdefault(instance, (model_id, class_zh))

    # 数据（每表时间轴 + 列值）
    tz_degraded_rows = 0
    per_sheet: dict[str, tuple[list[datetime], dict[str, list[Any]]]] = {}
    for sheet, cols in sheet_cols.items():
        ts: list[datetime] = []
        values: dict[str, list[Any]] = {
            f"{inst}.{prop}": [] for _, inst, _m, prop in cols
        }
        bool_props = {
            f"{inst}.{prop}"
            for _, inst, model_name, prop in cols
            if _tfom_dtype(registry, MODEL_ID_MAP[model_name], prop) == "boolean"
        }
        for row in sheets[sheet].data_rows:
            stamp = row[0] if row else None
            if stamp is None and all(v is None for v in row[1:]):
                continue
            if not isinstance(stamp, datetime):
                raise ValueError(
                    f"{sheet} 时间戳不是 Excel 序列号: {stamp!r}"
                )
            if stamp.tzinfo is not None:
                ts.append(stamp.astimezone(timezone.utc))
            else:
                # naive 序列号：按显式声明的 Asia/Shanghai 解释（降级）
                ts.append(stamp.replace(tzinfo=tz).astimezone(timezone.utc))
                tz_degraded_rows += 1
            for col_idx, instance, _m, prop in cols:
                vid = f"{instance}.{prop}"
                v = row[col_idx] if col_idx < len(row) else None
                if vid in bool_props and v is not None:
                    if v == 1:
                        v = True
                    elif v == 0:
                        v = False
                    # 其他取值（如 19）原样保留 → 管线报 TFDC-404
                values[vid].append(v)
        per_sheet[sheet] = (ts, values)

    # 并集时间轴 + 列对齐
    axis = sorted({t for ts, _ in per_sheet.values() for t in ts})
    # 下标必须按并集时间轴取，不能用表内行号：各表时间轴不一定相同
    # （某次导出里 冷冻水泵 表比其它表短 481 个时刻），用行号会让该表整列
    # 相对时间轴错位，甚至越界。
    axis_index = {t: i for i, t in enumerate(axis)}
    columns: dict[str, list[Any]] = {}
    for sheet, (ts, values) in per_sheet.items():
        for vid, vals in values.items():
            aligned: list[Any] = [None] * len(axis)
            target = columns.setdefault(vid, aligned)
            for i, t in enumerate(ts):
                target[axis_index[t]] = vals[i]

    # objects / variables
    object_records = [
        ObjectRecord(
            object_id=instance,
            object_model_id=model_id,
            object_name=class_zh,
        )
        for instance, (model_id, class_zh) in sorted(objects.items())
    ]
    variable_records: list[VariableRecord] = []
    for sheet, cols in sheet_cols.items():
        for _idx, instance, model_name, prop in cols:
            model_id = MODEL_ID_MAP[model_name]
            tfom_prop = registry.get(model_id).properties[prop]
            variable_records.append(
                VariableRecord(
                    variable_id=f"{instance}.{prop}",
                    object_id=instance,
                    property_code=prop,
                    unit=tfom_prop.unit,
                    dtype=tfom_prop.dtype,
                    role=tfom_prop.role,
                    source_kind="measured"
                    if tfom_prop.role != "derived" else "derived",
                    nullable=True,
                )
            )
    # 去重（同一变量可能被多张表引用——本工作簿不会，防御性处理）
    seen: set[str] = set()
    variable_records = [
        v for v in variable_records if not (v.variable_id in seen)
        and not seen.add(v.variable_id)
    ]

    manifest = TfdcManifest(
        contract="TFDC",
        contract_version="1.0",
        dataset_id=dataset_id,
        dataset_version=1,
        site_id=site_id,
        timezone=LEGACY_TIMEZONE,
        time_resolution=time_resolution,
        source_system="legacy_excel",
        description=f"旧格式工作簿转换: {source_name or '<内存表模型>'}",
    )
    dataset = TfdcDataset(
        manifest=manifest,
        objects=object_records,
        variables=variable_records,
    )

    pre_diagnostics = [
        Diagnostic(
            code="TFDC-502",
            level=Level.WARN,
            message=(
                "timestamp 为 Excel 序列号（naive），已按显式声明的 "
                f"manifest.timezone={LEGACY_TIMEZONE} 解释（降级行为）"
            ),
            location="data!timestamp",
            count=tz_degraded_rows,
        )
    ]
    degradations = [
        f"timestamp 为 naive Excel 序列号，按 Asia/Shanghai 显式声明解释"
        f"（{tz_degraded_rows} 行）"
    ]
    return LegacyConversion(
        dataset=dataset,
        timestamps=axis,
        columns=columns,
        property_mapping=property_mapping,
        pre_diagnostics=pre_diagnostics,
        degradations=degradations,
    )


def _tfom_dtype(registry: TfomRegistry, model_id: str, prop: str) -> str:
    model = registry.get(model_id)
    if model is None or prop not in model.properties:
        raise ValueError(f"TFOM 中未找到属性: {model_id}.{prop}")
    return model.properties[prop].dtype


def _parse_multi_header(
    header_rows: list[list[Any]],
) -> tuple[list[tuple[int, str, str, str]], dict[str, str]]:
    """解析 4 行多级表头 → ([(col_idx, instance, model_name, property_code)], 映射表)。

    r1–r3 因合并单元格只在前置列有值，需横向前向填充；
    r4 为空的列（实例合并范围之外的尾部列）跳过。
    映射表记录被改名的属性（如 `电流百分比` → `current_percent`）。
    """
    if len(header_rows) < 4:
        raise ValueError("旧格式工作簿必须有 4 行表头")
    r2, r3, r4 = header_rows[1], header_rows[2], header_rows[3]
    cols: list[tuple[int, str, str, str]] = []
    mapping: dict[str, str] = {}
    instance = model = None
    for j in range(1, max(len(r2), len(r4))):
        if j < len(r2) and r2[j] is not None:
            instance = str(r2[j])
        if j < len(r3) and r3[j] is not None:
            model = str(r3[j])
        raw_prop = r4[j] if j < len(r4) else None
        if raw_prop is None or instance is None or model is None:
            continue
        prop = PROPERTY_NAME_MAP.get(str(raw_prop), str(raw_prop))
        if prop != str(raw_prop):
            mapping[str(raw_prop)] = prop
        if not is_property_code(prop):
            raise ValueError(
                f"属性名无法映射为 property_code: {raw_prop!r}，"
                "请在 PROPERTY_NAME_MAP 登记"
            )
        cols.append((j, instance, model, prop))
    return cols, mapping
