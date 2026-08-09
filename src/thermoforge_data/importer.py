"""TFDC-XLSX 导入器与验证器（data-contract.md §4–6、conventions.md §7）。

管线：ZIP 层结构检查（implementation-notes §1.2）→ openpyxl 只读流式解析 →
清单/引用/单位/类型校验 → 时间轴与数值质量校验 → Arrow 内部表示。

- ERROR：中止，不产生 revision（`ImportResult.ok=False`，`table=None`）。
- REJECT：剔除相关单元格（置 null）并记录；单列剔除比例超阈值升级为 ERROR。
- WARN：记入诊断，导入继续。

所有诊断使用 `thermoforge_core.errors.Diagnostic`，同类按
(code, level, location) 聚合计数（conventions.md §7.1）。
"""

from __future__ import annotations

import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from zoneinfo import ZoneInfo

import pyarrow as pa
import yaml
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from thermoforge_core.contracts.tfom import ObjectModel
from thermoforge_core.contracts.tfdc import (
    BindingRecord,
    ObjectRecord,
    ParameterRecord,
    RelationRecord,
    TfdcDataset,
    TfdcManifest,
    VariableRecord,
)
from thermoforge_core.errors import Diagnostic, Level
from thermoforge_core.naming import (
    is_object_id,
    is_object_model_id,
    is_property_code,
    is_variable_id,
    split_variable_id,
)
from thermoforge_core.timeutil import (
    TIME_RESOLUTION_RE,
    is_iana_timezone,
    parse_time_resolution,
    parse_timestamp,
)
from thermoforge_core.units import (
    UnitError,
    conversion_factor,
    normalize_unit,
)

IMPORTER_VERSION = "thermoforge-data 0.1.0"

REQUIRED_SHEETS = ("manifest", "objects", "variables", "data")
KNOWN_SHEETS = frozenset(
    REQUIRED_SHEETS + ("parameters", "relations", "bindings", "events", "quality")
)

MAX_SHEET_ROWS = 1_048_576
MAX_SHEET_COLS = 16_384

_ARROW_TYPES = {
    "float": pa.float64(),
    "integer": pa.int64(),
    "boolean": pa.bool_(),
    "string": pa.string(),
}


# ---------------------------------------------------------------- 诊断收集


class DiagnosticSink:
    """按 (code, level, location) 聚合计数的诊断收集器（§7.1）。"""

    def __init__(self) -> None:
        self._items: dict[tuple[str, str, str | None], Diagnostic] = {}
        self._order: list[tuple[str, str, str | None]] = []

    def add(
        self,
        code: str,
        level: Level,
        message: str,
        location: str | None = None,
        count: int = 1,
    ) -> None:
        key = (code, level.value, location)
        if key in self._items:
            prev = self._items[key]
            self._items[key] = Diagnostic(
                code=prev.code,
                level=prev.level,
                message=prev.message,
                location=prev.location,
                count=prev.count + count,
            )
        else:
            self._items[key] = Diagnostic(
                code=code, level=level, message=message, location=location, count=count
            )
            self._order.append(key)

    def extend(self, diagnostics: Iterable[Diagnostic]) -> None:
        for d in diagnostics:
            self.add(d.code, d.level, d.message, d.location, d.count)

    @property
    def diagnostics(self) -> list[Diagnostic]:
        return [self._items[k] for k in self._order]

    def has_level(self, level: Level) -> bool:
        return any(d.level == level for d in self._items.values())


# ---------------------------------------------------------------- 选项与结果


@dataclass(frozen=True)
class ImportOptions:
    """导入阈值与开关。默认值标注 [草案]，需用真实数据标定（conventions §8 #7）。"""

    allow_naive_with_manifest_tz: bool = False  # §3.1 降级授权
    gap_factor: float = 3.0  # TFDC-505：间隔超过 N 倍采样周期
    drift_tolerance: float = 0.01  # TFDC-506：偏离声明周期的间隔比例上限
    missing_rate_warn: float = 0.2  # TFDC-607
    reject_escalation: float = 0.1  # 单列 REJECT 比例超过则升级为 ERROR
    derived_abs_tol: float = 1e-6  # TFDC-604 容差 max(abs, rel·|expected|)
    derived_rel_tol: float = 1e-6
    sentinel_text: tuple[str, ...] = ("NULL", "N/A", "NA", "--", "-", "#N/A")
    sentinel_numeric: tuple[float, ...] = (9999.0, -9999.0, 99999.0, 32767.0, 65535.0)
    supported_contract_major: int = 1


@dataclass
class ImportResult:
    """导入结果。`ok=False` 表示存在 ERROR 级诊断，不得落 vault。"""

    ok: bool
    dataset: TfdcDataset | None
    table: pa.Table | None  # data 表：timestamp[us, UTC] + 变量列
    diagnostics: list[Diagnostic]
    degradations: list[str]  # 降级行为记录（如 naive 时间戳按 manifest 时区解释）
    source_path: Path | None
    importer_version: str = IMPORTER_VERSION

    @property
    def has_errors(self) -> bool:
        return any(d.level == Level.ERROR for d in self.diagnostics)


# ---------------------------------------------------------------- TFOM 注册表


class TfomRegistry:
    """物模型注册表：`object_model_id`（name.vN）→ `ObjectModel`。"""

    def __init__(self, models: Mapping[str, ObjectModel]):
        self._models = dict(models)

    @classmethod
    def from_dir(cls, directory: str | Path) -> "TfomRegistry":
        models: dict[str, ObjectModel] = {}
        for path in sorted(Path(directory).glob("*.yaml")):
            with open(path, encoding="utf-8") as fp:
                model = ObjectModel.model_validate(yaml.safe_load(fp))
            models[model.object_model_id] = model
        return cls(models)

    def get(self, object_model_id: str) -> ObjectModel | None:
        return self._models.get(object_model_id)

    def as_dict(self) -> dict[str, ObjectModel]:
        """注册表内容的只读拷贝（合并/扩展注册表时使用）。"""
        return dict(self._models)

    def __contains__(self, object_model_id: str) -> bool:
        return object_model_id in self._models


def default_registry() -> TfomRegistry:
    """仓库内置示例注册表（contracts/tfom/examples/）。"""
    examples = (
        Path(__file__).resolve().parents[2] / "contracts" / "tfom" / "examples"
    )
    return TfomRegistry.from_dir(examples)


# ---------------------------------------------------------------- ZIP 结构检查


def _sheet_xml_map(zf: zipfile.ZipFile) -> dict[str, str]:
    """工作表名 → xl/worksheets/sheetN.xml 路径。"""
    wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
    rels_xml = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rels = dict(
        re.findall(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels_xml)
    )
    # Target 在前 Id 在后的写法也兼容
    for m in re.finditer(r"<Relationship\b[^>]*>", rels_xml):
        tag = m.group(0)
        rid = re.search(r'Id="([^"]+)"', tag)
        target = re.search(r'Target="([^"]+)"', tag)
        if rid and target:
            rels[rid.group(1)] = target.group(1)
    mapping: dict[str, str] = {}
    for m in re.finditer(r"<sheet\b[^>]*>", wb_xml):
        tag = m.group(0)
        name = re.search(r'name="([^"]+)"', tag)
        rid = re.search(r'r:id="([^"]+)"', tag)
        if name and rid and rid.group(1) in rels:
            target = rels[rid.group(1)].lstrip("/")
            if not target.startswith("xl/"):
                target = "xl/" + target
            mapping[name.group(1)] = target
    return mapping


def _structural_checks(path: Path, sink: DiagnosticSink) -> bool:
    """ZIP 层结构检查（implementation-notes §1.2）。返回 False 表示致命错误。"""
    if path.suffix.lower() != ".xlsx":
        sink.add(
            "TFDC-101", Level.ERROR,
            f"非 .xlsx 文件: {path.name}", location=str(path),
        )
        return False
    if not zipfile.is_zipfile(path):
        sink.add(
            "TFDC-101", Level.ERROR,
            f"不是合法的 xlsx（ZIP 容器）: {path.name}", location=str(path),
        )
        return False
    fatal = False
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        if "xl/vbaProject.bin" in names:
            sink.add("TFDC-102", Level.ERROR, "工作簿含宏或 VBA 工程",
                     location=str(path))
            fatal = True
        if any(n.startswith("xl/externalLinks/") for n in names):
            sink.add("TFDC-109", Level.ERROR, "存在外部工作簿引用",
                     location=str(path))
            fatal = True
        wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
        if re.search(r'date1904="(1|true)"', wb_xml):
            sink.add("TFDC-202", Level.ERROR, "禁止 1904 日期系统（conventions §3.1）",
                     location="xl/workbook.xml")
            fatal = True
        sheet_map = _sheet_xml_map(zf)
        for sheet_name, xml_path in sheet_map.items():
            xml = zf.read(xml_path).decode("utf-8")
            n_merged = len(re.findall(r"<mergeCell[ />]", xml))
            if n_merged:
                sink.add("TFDC-105", Level.ERROR,
                         f"存在 {n_merged} 个合并单元格",
                         location=sheet_name, count=n_merged)
                fatal = True
            if sheet_name == "data":
                n_formula = len(re.findall(r"<f[ />]", xml))
                if n_formula:
                    sink.add("TFDC-108", Level.ERROR,
                             f"data 中存在 {n_formula} 个公式",
                             location="data", count=n_formula)
                    fatal = True
            dim = re.search(r'<dimension ref="([^"]+)"', xml)
            if dim:
                last = dim.group(1).split(":")[-1]
                m = re.match(r"([A-Z]+)([0-9]+)", last)
                if m:
                    col_s, row_s = m.group(1), int(m.group(2))
                    col_n = 0
                    for ch in col_s:
                        col_n = col_n * 26 + (ord(ch) - 64)
                    if row_s > MAX_SHEET_ROWS or col_n > MAX_SHEET_COLS:
                        sink.add("TFDC-110", Level.ERROR,
                                 f"超出 xlsx 单表上限: {dim.group(1)}",
                                 location=sheet_name)
                        fatal = True
    return not fatal


# ---------------------------------------------------------------- 表解析


def _check_header(header: Sequence[Any], sheet: str, sink: DiagnosticSink) -> bool:
    """表头空列名（106）/重复（107）检查。"""
    ok = True
    seen: set[str] = set()
    for idx, cell in enumerate(header):
        if cell is None or (isinstance(cell, str) and not cell.strip()):
            sink.add("TFDC-106", Level.ERROR,
                     f"表头第 {idx + 1} 列为空", location=sheet)
            ok = False
            continue
        name = str(cell)
        if name in seen:
            sink.add("TFDC-107", Level.ERROR,
                     f"表头列名重复: {name!r}", location=sheet)
            ok = False
        seen.add(name)
    return ok


def _read_sheet_rows(ws) -> tuple[list[Any], list[tuple[Any, ...]]]:
    """只读流式读取一张表，返回 (header, rows)。"""
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = list(next(rows_iter))
    except StopIteration:
        return [], []
    # 去掉全空的尾部列
    while header and header[-1] is None:
        header.pop()
    return header, [tuple(r[: len(header)]) for r in rows_iter]


# ---------------------------------------------------------------- 表达式求值（TFDC-604）

import ast  # noqa: E402

_ALLOWED_BINOPS = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b,
                   ast.Mult: lambda a, b: a * b, ast.Div: lambda a, b: a / b}


def _eval_expression(expr: str, env: Mapping[str, float]) -> float:
    """求值 TFOM 派生表达式：仅允许 + - * /、括号、数值与属性名。"""

    def walk(node: ast.AST) -> float:
        if isinstance(node, ast.Expression):
            return walk(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return float(node.value)
        if isinstance(node, ast.Name):
            if node.id not in env:
                raise ValueError(f"表达式引用未知属性: {node.id!r}")
            return env[node.id]
        if isinstance(node, ast.BinOp) and type(node.op) in _ALLOWED_BINOPS:
            return _ALLOWED_BINOPS[type(node.op)](walk(node.left), walk(node.right))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            v = walk(node.operand)
            return v if isinstance(node.op, ast.UAdd) else -v
        raise ValueError(f"表达式含不支持的语法: {ast.dump(node)}")

    return walk(ast.parse(expr, mode="eval"))


def expression_dependencies(expr: str) -> set[str]:
    """表达式引用的属性名集合。"""
    return {
        node.id
        for node in ast.walk(ast.parse(expr, mode="eval"))
        if isinstance(node, ast.Name)
    }


# ---------------------------------------------------------------- 主入口


def import_xlsx(
    path: str | Path,
    *,
    options: ImportOptions | None = None,
    registry: TfomRegistry | None = None,
) -> ImportResult:
    """导入 TFDC-XLSX 工作簿，执行全部契约校验。"""
    options = options or ImportOptions()
    registry = registry or default_registry()
    path = Path(path)
    sink = DiagnosticSink()
    degradations: list[str] = []

    if not _structural_checks(path, sink):
        return ImportResult(False, None, None, sink.diagnostics, degradations, path)

    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet_names = list(wb.sheetnames)
        missing = [s for s in REQUIRED_SHEETS if s not in sheet_names]
        if missing:
            for name in missing:
                sink.add("TFDC-103", Level.ERROR, f"缺少必需工作表: {name}",
                         location=name)
            return ImportResult(False, None, None, sink.diagnostics, degradations, path)
        for name in sheet_names:
            if name not in KNOWN_SHEETS:
                sink.add("TFDC-104", Level.WARN, f"契约未定义的工作表: {name}",
                         location=name)

        # ---- manifest ----
        manifest = _parse_manifest(wb["manifest"], sink)
        if manifest is None:
            return ImportResult(False, None, None, sink.diagnostics, degradations, path)

        # ---- objects / variables / 可选表 ----
        objects = _parse_objects(wb["objects"], sink, registry)
        variables = _parse_variables(wb["variables"], sink)
        parameters = _parse_optional(
            wb, "parameters", PARAMETERS_COLUMNS, sink) if "parameters" in sheet_names else []
        relations = _parse_optional(
            wb, "relations", RELATIONS_COLUMNS, sink) if "relations" in sheet_names else []
        bindings = _parse_optional(
            wb, "bindings", BINDINGS_COLUMNS, sink) if "bindings" in sheet_names else []

        # ---- data ----
        parsed_data = _parse_data_sheet(wb["data"], sink)
    finally:
        wb.close()

    if sink.has_level(Level.ERROR) or objects is None or variables is None or parsed_data is None:
        return ImportResult(False, None, None, sink.diagnostics, degradations, path)

    raw_timestamps, raw_columns, data_column_order = parsed_data

    dataset = _build_dataset(
        manifest, objects, variables, parameters, relations, bindings, sink
    )
    if dataset is None:
        return ImportResult(False, None, None, sink.diagnostics, degradations, path)

    return _finalize(
        dataset, raw_timestamps, raw_columns, data_column_order,
        sink=sink, options=options, registry=registry,
        source_path=path, degradations=degradations,
    )


def import_parsed(
    dataset: TfdcDataset,
    raw_timestamps: Sequence[Any],
    raw_columns: Mapping[str, Sequence[Any]],
    *,
    options: ImportOptions | None = None,
    registry: TfomRegistry | None = None,
    source_path: str | Path | None = None,
    pre_diagnostics: Iterable[Diagnostic] = (),
    degradations: Iterable[str] = (),
) -> ImportResult:
    """标准管线的内存入口：旧格式适配器等预处理步骤复用全部校验。"""
    options = options or ImportOptions()
    registry = registry or default_registry()
    sink = DiagnosticSink()
    sink.extend(pre_diagnostics)
    column_order = list(raw_columns.keys())
    return _finalize(
        dataset, list(raw_timestamps), {k: list(v) for k, v in raw_columns.items()},
        column_order,
        sink=sink, options=options, registry=registry,
        source_path=Path(source_path) if source_path else None,
        degradations=list(degradations),
    )


# ---------------------------------------------------------------- manifest


def _parse_manifest(ws, sink: DiagnosticSink) -> TfdcManifest | None:
    header, rows = _read_sheet_rows(ws)
    if not header or not _check_header(header, "manifest", sink):
        sink.add("TFDC-202", Level.ERROR, "manifest 表头非法", location="manifest")
        return None
    kv: dict[str, Any] = {}
    key_idx = header.index("key") if "key" in header else 0
    val_idx = header.index("value") if "value" in header else 1
    for row in rows:
        if row[key_idx] is None:
            continue
        kv[str(row[key_idx])] = row[val_idx]

    required = ("contract", "contract_version", "dataset_id", "dataset_version",
                "site_id", "timezone", "time_resolution")
    missing = [k for k in required if k not in kv or kv[k] is None]
    if missing:
        for k in missing:
            sink.add("TFDC-201", Level.ERROR, f"manifest 缺少必需 key: {k}",
                     location="manifest")
        return None

    if str(kv["contract"]) != "TFDC":
        sink.add("TFDC-202", Level.ERROR,
                 f"contract 必须为 TFDC: {kv['contract']!r}", location="manifest")
        return None
    version = str(kv["contract_version"])
    if not re.fullmatch(r"[0-9]+\.[0-9]+", version):
        sink.add("TFDC-202", Level.ERROR,
                 f"contract_version 格式非法: {version!r}", location="manifest")
        return None
    major = int(version.split(".", 1)[0])
    if major > 1:
        sink.add("TFDC-203", Level.ERROR,
                 f"contract_version MAJOR 超出支持范围: {version}", location="manifest")
        return None
    if not is_iana_timezone(str(kv["timezone"])):
        sink.add("TFDC-204", Level.ERROR,
                 f"非 IANA 时区名: {kv['timezone']!r}", location="manifest")
        return None
    if not TIME_RESOLUTION_RE.fullmatch(str(kv["time_resolution"])):
        sink.add("TFDC-205", Level.ERROR,
                 f"非法 time_resolution: {kv['time_resolution']!r}",
                 location="manifest")
        return None
    if kv.get("created_at") is not None and not isinstance(kv["created_at"], datetime):
        try:
            kv["created_at"] = parse_timestamp(str(kv["created_at"]))
        except ValueError as exc:
            sink.add("TFDC-202", Level.ERROR, f"created_at 非法: {exc}",
                     location="manifest")
            return None
    try:
        return TfdcManifest.model_validate(kv)
    except ValueError as exc:
        sink.add("TFDC-202", Level.ERROR, f"manifest 校验失败: {exc}",
                 location="manifest")
        return None


# ---------------------------------------------------------------- objects / variables


def _parse_objects(ws, sink: DiagnosticSink,
                   registry: TfomRegistry) -> list[ObjectRecord] | None:
    header, rows = _read_sheet_rows(ws)
    if not header or not _check_header(header, "objects", sink):
        return None
    header = [str(h) for h in header]
    for col in ("object_id", "object_model_id"):
        if col not in header:
            sink.add("TFDC-202", Level.ERROR, f"objects 表缺少必需列: {col}",
                     location="objects")
            return None
    records: list[ObjectRecord] = []
    seen: set[str] = set()
    ok = True
    for i, row in enumerate(rows, start=2):
        rec = {h: row[j] for j, h in enumerate(header)}
        rec = {k: v for k, v in rec.items() if v is not None}
        object_id = str(rec.get("object_id", ""))
        model_id = str(rec.get("object_model_id", ""))
        if not is_object_id(object_id):
            sink.add("TFDC-304", Level.ERROR,
                     f"非法 object_id: {object_id!r}", location=f"objects!row{i}")
            ok = False
            continue
        if object_id in seen:
            sink.add("TFDC-307", Level.ERROR,
                     f"object_id 重复: {object_id!r}", location=f"objects!row{i}")
            ok = False
            continue
        seen.add(object_id)
        if not is_object_model_id(model_id):
            sink.add("TFDC-304", Level.ERROR,
                     f"非法 object_model_id: {model_id!r}", location=f"objects!row{i}")
            ok = False
            continue
        if model_id not in registry:
            sink.add("TFDC-302", Level.ERROR,
                     f"object_model_id 未在 TFOM Registry 注册: {model_id!r}",
                     location=f"objects!row{i}")
            ok = False
            continue
        try:
            records.append(ObjectRecord.model_validate(rec))
        except ValueError as exc:
            sink.add("TFDC-202", Level.ERROR, f"objects 行校验失败: {exc}",
                     location=f"objects!row{i}")
            ok = False
    return records if ok else None


VARIABLES_COLUMNS = ("variable_id", "object_id", "property_code", "unit",
                     "dtype", "role", "source_kind")


def _parse_variables(ws, sink: DiagnosticSink) -> list[VariableRecord] | None:
    header, rows = _read_sheet_rows(ws)
    if not header or not _check_header(header, "variables", sink):
        return None
    header = [str(h) for h in header]
    for col in VARIABLES_COLUMNS:
        if col not in header:
            sink.add("TFDC-202", Level.ERROR, f"variables 表缺少必需列: {col}",
                     location="variables")
            return None
    records: list[VariableRecord] = []
    seen: set[str] = set()
    ok = True
    for i, row in enumerate(rows, start=2):
        rec = {h: row[j] for j, h in enumerate(header) if j < len(row)}
        rec = {k: v for k, v in rec.items() if v is not None}
        variable_id = str(rec.get("variable_id", ""))
        loc = f"variables!row{i}"
        if variable_id in seen:
            sink.add("TFDC-307", Level.ERROR,
                     f"variable_id 重复声明: {variable_id!r}", location=loc)
            ok = False
            continue
        seen.add(variable_id)
        if not is_variable_id(variable_id):
            sink.add("TFDC-304", Level.ERROR,
                     f"非法 variable_id: {variable_id!r}", location=loc)
            ok = False
            continue
        object_id, property_code = split_variable_id(variable_id)
        if (str(rec.get("object_id", "")) != object_id
                or str(rec.get("property_code", "")) != property_code):
            sink.add("TFDC-304", Level.ERROR,
                     f"variable_id 与 object_id/property_code 不一致: {variable_id!r}",
                     location=loc)
            ok = False
            continue
        try:
            rec["unit"] = normalize_unit(str(rec.get("unit", "")))
        except UnitError:
            sink.add("TFDC-401", Level.ERROR,
                     f"未登记的单位字符串: {rec.get('unit')!r}", location=loc)
            ok = False
            continue
        try:
            records.append(VariableRecord.model_validate(rec))
        except ValueError as exc:
            sink.add("TFDC-202", Level.ERROR, f"variables 行校验失败: {exc}",
                     location=loc)
            ok = False
    return records if ok else None


PARAMETERS_COLUMNS = ("object_id", "parameter_code", "value", "unit")
RELATIONS_COLUMNS = ("from_object", "relation", "to_object")
BINDINGS_COLUMNS = ("variable_id", "adapter", "source_ref")

_OPTIONAL_MODELS = {
    "parameters": (ParameterRecord, PARAMETERS_COLUMNS),
    "relations": (RelationRecord, RELATIONS_COLUMNS),
    "bindings": (BindingRecord, BINDINGS_COLUMNS),
}


def _parse_optional(wb, sheet: str, columns, sink: DiagnosticSink) -> list:
    header, rows = _read_sheet_rows(wb[sheet])
    if not header:
        return []
    _check_header(header, sheet, sink)
    header = [str(h) for h in header]
    model, required = _OPTIONAL_MODELS[sheet]
    records = []
    for i, row in enumerate(rows, start=2):
        rec = {h: row[j] for j, h in enumerate(header) if j < len(row)}
        rec = {k: v for k, v in rec.items() if v is not None}
        if not rec:
            continue
        try:
            records.append(model.model_validate(rec))
        except ValueError as exc:
            sink.add("TFDC-202", Level.ERROR, f"{sheet} 行校验失败: {exc}",
                     location=f"{sheet}!row{i}")
    return records


# ---------------------------------------------------------------- data 表解析


def _parse_data_sheet(
    ws, sink: DiagnosticSink
) -> tuple[list[Any], dict[str, list[Any]], list[str]] | None:
    """流式解析 data 表，返回 (原始 timestamp 列, 变量列, 列序)。"""
    header, rows = _read_sheet_rows(ws)
    if not header:
        sink.add("TFDC-103", Level.ERROR, "data 表为空", location="data")
        return None
    if not _check_header(header, "data", sink):
        return None
    header = [str(h) for h in header]
    if header[0] != "timestamp":
        sink.add("TFDC-202", Level.ERROR,
                 f"data 第一列必须为 timestamp: {header[0]!r}", location="data")
        return None
    columns = header[1:]
    raw_columns: dict[str, list[Any]] = {c: [] for c in columns}
    raw_timestamps: list[Any] = []
    for row in rows:
        if all(v is None for v in row):
            continue  # 尾部空行
        raw_timestamps.append(row[0])
        for j, col in enumerate(columns):
            raw_columns[col].append(row[j + 1] if j + 1 < len(row) else None)
    return raw_timestamps, raw_columns, columns


def _build_dataset(
    manifest: TfdcManifest,
    objects: list[ObjectRecord],
    variables: list[VariableRecord],
    parameters: list,
    relations: list,
    bindings: list,
    sink: DiagnosticSink,
) -> TfdcDataset | None:
    try:
        return TfdcDataset(
            manifest=manifest, objects=objects, variables=variables,
            parameters=parameters, relations=relations, bindings=bindings,
        )
    except ValueError as exc:
        sink.add("TFDC-301", Level.ERROR, f"跨表引用校验失败: {exc}",
                 location="objects/variables")
        return None


# ----------------------------------------------------------------  finalize：语义与质量校验


def _finalize(
    dataset: TfdcDataset,
    raw_timestamps: list[Any],
    raw_columns: dict[str, list[Any]],
    column_order: list[str],
    *,
    sink: DiagnosticSink,
    options: ImportOptions,
    registry: TfomRegistry,
    source_path: Path | None,
    degradations: list[str],
) -> ImportResult:
    manifest = dataset.manifest
    n_rows = len(raw_timestamps)
    variables = {v.variable_id: v for v in dataset.variables}
    object_model = {o.object_id: o.object_model_id for o in dataset.objects}

    # ---- data 列 ↔ variables 声明 ----
    for col in column_order:
        if not is_variable_id(col):
            sink.add("TFDC-304", Level.ERROR, f"非法 variable_id: {col!r}",
                     location=f"data!{col}")
        elif col not in variables:
            sink.add("TFDC-305", Level.ERROR, f"data 列未在 variables 声明: {col}",
                     location=f"data!{col}")
    for var in dataset.variables:
        if var.variable_id not in raw_columns:
            level = Level.WARN if var.nullable else Level.ERROR
            sink.add("TFDC-306", level,
                     f"variables 已声明但 data 无对应列: {var.variable_id}",
                     location=f"variables!{var.variable_id}")
            raw_columns[var.variable_id] = [None] * n_rows
            column_order.append(var.variable_id)
    if sink.has_level(Level.ERROR):
        return ImportResult(False, dataset, None, sink.diagnostics, degradations,
                            source_path)

    # ---- 引用与单位校验（301/303/401/402/403）----
    conversions: dict[str, tuple[float, float]] = {}
    ranges: dict[str, tuple[float | None, float | None]] = {}
    for var in dataset.variables:
        loc = f"variables!{var.variable_id}"
        if var.object_id not in object_model:
            sink.add("TFDC-301", Level.ERROR,
                     f"引用 objects 中不存在的对象: {var.object_id}", location=loc)
            continue
        model = registry.get(object_model[var.object_id])
        if model is None:
            sink.add("TFDC-302", Level.ERROR,
                     f"object_model_id 未注册: {object_model[var.object_id]}",
                     location=loc)
            continue
        prop = model.properties.get(var.property_code)
        if prop is None:
            sink.add("TFDC-303", Level.ERROR,
                     f"property_code 不属于 {model.object_model_id}: "
                     f"{var.property_code!r}", location=loc)
            continue
        if var.unit != prop.unit:
            try:
                factor = conversion_factor(var.unit, prop.unit, prop.quantity_kind)
            except UnitError as exc:
                sink.add(exc.code, Level.ERROR, str(exc), location=loc)
                continue
            if factor != (1.0, 0.0):
                conversions[var.variable_id] = factor
        # 合法范围：variables 声明优先，缺省回退到 TFOM 属性（TFDC-601）
        ranges[var.variable_id] = (
            var.min_value if var.min_value is not None else prop.min_value,
            var.max_value if var.max_value is not None else prop.max_value,
        )
    if sink.has_level(Level.ERROR):
        return ImportResult(False, dataset, None, sink.diagnostics, degradations,
                            source_path)

    # ---- 时间戳解析 ----
    tz = ZoneInfo(manifest.timezone)
    timestamps = _convert_timestamps(
        raw_timestamps, tz, options, sink, degradations
    )
    if timestamps is None or sink.has_level(Level.ERROR):
        return ImportResult(False, dataset, None, sink.diagnostics, degradations,
                            source_path)

    # ---- 单元格落地：dtype / 哨兵 / inf / 范围 ----
    columns: dict[str, list[Any]] = {}
    for col in column_order:
        var = variables[col]
        values = _land_column(var, raw_columns[col], ranges.get(col, (None, None)),
                              options, sink)
        if col in conversions and var.dtype in ("float", "integer"):
            scale, offset = conversions[col]
            values = [None if v is None else v * scale + offset for v in values]
        columns[col] = values
    # REJECT 升级检查：单列某类 REJECT 超阈值即升级为同级 ERROR（§7.1）
    if n_rows:
        reject_counts: dict[tuple[str, str], int] = {}
        for d in sink.diagnostics:
            if d.level == Level.REJECT and d.location:
                key = (d.code, d.location)
                reject_counts[key] = reject_counts.get(key, 0) + d.count
        for (code, loc), count in reject_counts.items():
            if count / n_rows > options.reject_escalation:
                sink.add(code, Level.ERROR,
                         f"列剔除比例 {count / n_rows:.1%} 超过阈值 "
                         f"{options.reject_escalation:.0%}，升级为 ERROR",
                         location=loc, count=count)
    if sink.has_level(Level.ERROR):
        return ImportResult(False, dataset, None, sink.diagnostics, degradations,
                            source_path)

    # ---- 时间轴：重复 / 乱序 / 空洞 / 漂移 ----
    _check_timeline(timestamps, columns, column_order, manifest,
                    options, sink)
    if sink.has_level(Level.ERROR):
        return ImportResult(False, dataset, None, sink.diagnostics, degradations,
                            source_path)

    # ---- 数值质量：恒定 / 分布 / 缺失率 ----
    for col in column_order:
        _check_series_quality(variables[col], columns[col], options, sink)

    # ---- 派生量一致性（604）----
    _check_derived(dataset, registry, object_model, columns, options, sink)
    if sink.has_level(Level.ERROR):
        return ImportResult(False, dataset, None, sink.diagnostics, degradations,
                            source_path)

    # ---- Arrow 内部表示 ----
    table = _build_arrow(timestamps, columns, column_order, variables)
    return ImportResult(True, dataset, table, sink.diagnostics, degradations,
                        source_path)


def _convert_timestamps(
    raw: list[Any],
    tz: ZoneInfo,
    options: ImportOptions,
    sink: DiagnosticSink,
    degradations: list[str],
) -> list[datetime] | None:
    out: list[datetime] = []
    naive_count = 0
    for i, value in enumerate(raw):
        loc = f"data!A{i + 2}"
        if value is None:
            sink.add("TFDC-501", Level.ERROR, "timestamp 为空", location=loc)
            out.append(datetime.now(timezone.utc))  # 占位，导入将中止
            continue
        if isinstance(value, str):
            try:
                out.append(parse_timestamp(value))
                continue
            except ValueError:
                pass  # 落入下面的分类
            try:
                dt = datetime.fromisoformat(value)
            except ValueError:
                sink.add("TFDC-501", Level.ERROR,
                         f"无法解析的时间戳: {value!r}", location=loc)
                out.append(datetime.now(timezone.utc))
                continue
        elif isinstance(value, datetime):
            dt = value
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            dt = from_excel(value)
        else:
            sink.add("TFDC-501", Level.ERROR,
                     f"无法解析的时间戳类型: {type(value).__name__}", location=loc)
            out.append(datetime.now(timezone.utc))
            continue
        if dt.tzinfo is not None:
            out.append(dt.astimezone(timezone.utc))
            continue
        # naive：Excel 序列号或无偏移文本（§3.1）
        naive_count += 1
        if not options.allow_naive_with_manifest_tz:
            sink.add("TFDC-502", Level.ERROR,
                     "timestamp 为 Excel 序列号/无时区文本，且未授权按 manifest "
                     "时区降级", location="data!timestamp")
            out.append(datetime.now(timezone.utc))
            continue
        naive_utc, status = _localize_naive(dt, tz)
        if status == "nonexistent":
            sink.add("TFDC-501", Level.ERROR,
                     f"本地时刻不存在（DST 跳变）: {dt!r}", location=loc)
        elif status == "ambiguous":
            sink.add("TFDC-507", Level.ERROR,
                     f"DST 折返导致的重复本地时刻: {dt!r}", location=loc)
        out.append(naive_utc)
    if naive_count and options.allow_naive_with_manifest_tz:
        degradations.append(
            f"timestamp 为 naive（Excel 序列号/无偏移文本），按 "
            f"manifest.timezone 显式声明解释（{naive_count} 行）"
        )
        sink.add("TFDC-502", Level.WARN,
                 f"naive 时间戳已按 manifest.timezone 解释（降级，已授权）",
                 location="data!timestamp", count=naive_count)
    if sink.has_level(Level.ERROR):
        return None
    return out


def _localize_naive(dt: datetime, tz: ZoneInfo) -> tuple[datetime, str]:
    """naive 本地时间 → UTC；返回 (utc, ok|nonexistent|ambiguous)。"""
    aware0 = dt.replace(tzinfo=tz, fold=0)
    aware1 = dt.replace(tzinfo=tz, fold=1)
    utc0 = aware0.astimezone(timezone.utc)
    utc1 = aware1.astimezone(timezone.utc)
    if utc0 != utc1:
        return utc0, "ambiguous"
    if utc0.astimezone(tz).replace(tzinfo=None) != dt:
        return utc0, "nonexistent"
    return utc0, "ok"


def _land_column(
    var: VariableRecord,
    raw: Sequence[Any],
    value_range: tuple[float | None, float | None],
    options: ImportOptions,
    sink: DiagnosticSink,
) -> list[Any]:
    """按声明 dtype 落地一列；返回落地后的值（违规单元格置 null）。"""
    loc = f"data!{var.variable_id}"
    out: list[Any] = []
    for value in raw:
        out.append(_land_cell(var, value, value_range, options, sink, loc))
    return out


def _land_cell(
    var: VariableRecord,
    value: Any,
    value_range: tuple[float | None, float | None],
    options: ImportOptions,
    sink: DiagnosticSink,
    loc: str,
) -> Any:
    if value is None:
        if not var.nullable:
            sink.add("TFDC-602", Level.ERROR,
                     "nullable=false 的变量出现空值", location=loc)
        return None
    if isinstance(value, str):
        stripped = value.strip().strip("　").strip()
        if not stripped:
            # 空白字符串：数值列等同缺失，字符串列报 603（§4.2）
            if var.dtype == "string":
                sink.add("TFDC-603", Level.ERROR,
                         f"字符串列出现空白字符串: {value!r}", location=loc)
            elif not var.nullable:
                sink.add("TFDC-602", Level.ERROR,
                         "nullable=false 的变量出现空值", location=loc)
            return None
        if stripped in options.sentinel_text:
            sink.add("TFDC-603", Level.ERROR,
                     f"检测到哨兵值: {value!r}", location=loc)
            return None
        value = stripped

    dtype = var.dtype
    if dtype == "float":
        v = _coerce_float(value, sink, loc)
    elif dtype == "integer":
        v = _coerce_int(value, sink, loc)
    elif dtype == "boolean":
        v = _coerce_bool(value, sink, loc)
    else:  # string
        if not isinstance(value, str):
            sink.add("TFDC-404", Level.REJECT,
                     f"string 列出现非字符串值: {value!r}", location=loc)
            return None
        v = value
    if v is None:
        return None
    if dtype in ("float", "integer"):
        fv = float(v)
        if math.isinf(fv):
            sink.add("TFDC-602", Level.ERROR, "检测到 inf/-inf", location=loc)
            return None
        if math.isnan(fv):
            sink.add("TFDC-602", Level.ERROR, "检测到 NaN（导入层不接受 NaN）",
                     location=loc)
            return None
        if fv in options.sentinel_numeric:
            sink.add("TFDC-603", Level.ERROR,
                     f"检测到哨兵值: {v!r}", location=loc)
            return None
        if fv == 0.0:
            fv = 0.0  # -0.0 规范化（§4.3）
            v = fv if dtype == "float" else int(fv)
        min_v, max_v = value_range
        if (min_v is not None and fv < min_v) or (max_v is not None and fv > max_v):
            sink.add("TFDC-601", Level.REJECT,
                     f"超出合法范围 [{min_v}, {max_v}]: {fv!r}", location=loc)
            return None
    return v


def _coerce_float(value: Any, sink: DiagnosticSink, loc: str) -> float | None:
    if isinstance(value, bool):
        sink.add("TFDC-404", Level.REJECT, f"float 列出现布尔值: {value!r}",
                 location=loc)
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value)
        except ValueError:
            pass
    sink.add("TFDC-404", Level.REJECT,
             f"float 列出现非数值: {value!r}", location=loc)
    return None


def _coerce_int(value: Any, sink: DiagnosticSink, loc: str) -> int | None:
    if isinstance(value, bool):
        sink.add("TFDC-404", Level.REJECT, f"integer 列出现布尔值: {value!r}",
                 location=loc)
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer() and math.isfinite(value):
            return int(value)
        sink.add("TFDC-404", Level.REJECT,
                 f"integer 列出现非整数值: {value!r}", location=loc)
        return None
    if isinstance(value, str):
        try:
            return int(value)
        except ValueError:
            pass
    sink.add("TFDC-404", Level.REJECT,
             f"integer 列出现非整数: {value!r}", location=loc)
    return None


def _coerce_bool(value: Any, sink: DiagnosticSink, loc: str) -> bool | None:
    # §4.1：只接受 TRUE/FALSE，不接受 1/0/是/否
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value in ("TRUE", "FALSE"):
        return value == "TRUE"
    sink.add("TFDC-404", Level.REJECT,
             f"boolean 列只接受 TRUE/FALSE: {value!r}", location=loc)
    return None


def _check_timeline(
    timestamps: list[datetime],
    columns: dict[str, list[Any]],
    column_order: list[str],
    manifest: TfdcManifest,
    options: ImportOptions,
    sink: DiagnosticSink,
) -> None:
    n = len(timestamps)
    # 重复时间戳（503）
    first_seen: dict[datetime, int] = {}
    dup_identical = 0
    dup_conflict = 0
    drop_rows: set[int] = set()
    for i, ts in enumerate(timestamps):
        if ts not in first_seen:
            first_seen[ts] = i
            continue
        j = first_seen[ts]
        identical = all(columns[c][i] == columns[c][j] for c in column_order)
        if identical:
            dup_identical += 1
            drop_rows.add(i)  # 去重保留首条（§3.5）
        else:
            dup_conflict += 1
    if dup_identical:
        sink.add("TFDC-503", Level.WARN,
                 "重复时间戳且各列值一致，已去重保留首条",
                 location="data!timestamp", count=dup_identical)
    if dup_conflict:
        sink.add("TFDC-503", Level.ERROR,
                 "重复时间戳且值不一致，中止导入（不得自动取均值或保留首条）",
                 location="data!timestamp", count=dup_conflict)
        return
    # 乱序（504）
    out_of_order = sum(
        1 for a, b in zip(timestamps, timestamps[1:]) if b < a
    )
    if out_of_order:
        sink.add("TFDC-504", Level.WARN,
                 "时间戳非单调递增，已按时间排序",
                 location="data!timestamp", count=out_of_order)
    # 去重 + 排序
    keep = sorted((i for i in range(n) if i not in drop_rows),
                  key=lambda i: timestamps[i])
    timestamps[:] = [timestamps[i] for i in keep]
    for col in column_order:
        columns[col] = [columns[col][i] for i in keep]
    # 空洞（505）与漂移（506）
    resolution_s = parse_time_resolution(manifest.time_resolution)
    diffs = [
        (b - a).total_seconds()
        for a, b in zip(timestamps, timestamps[1:])
    ]
    gaps = sum(1 for d in diffs if d > options.gap_factor * resolution_s)
    if gaps:
        sink.add("TFDC-505", Level.WARN,
                 f"存在超过 {options.gap_factor:g} 倍采样周期的空洞（不填充）",
                 location="data!timestamp", count=gaps)
    if diffs:
        drift = sum(1 for d in diffs if d != resolution_s) / len(diffs)
        if drift > options.drift_tolerance:
            sink.add("TFDC-506", Level.WARN,
                     f"{drift:.1%} 的采样间隔偏离声明值 "
                     f"{manifest.time_resolution}",
                     location="data!timestamp")


def _check_series_quality(
    var: VariableRecord,
    values: Sequence[Any],
    options: ImportOptions,
    sink: DiagnosticSink,
) -> None:
    loc = f"data!{var.variable_id}"
    n = len(values)
    non_null = [v for v in values if v is not None]
    missing = n - len(non_null)
    if n and missing / n > options.missing_rate_warn:
        sink.add("TFDC-607", Level.WARN,
                 f"缺失率 {missing / n:.1%} 超过阈值 "
                 f"{options.missing_rate_warn:.0%}",
                 location=loc, count=missing)
    if not non_null:
        return
    if var.dtype in ("float", "integer"):
        nums = [float(v) for v in non_null]
        if all(v == nums[0] for v in nums):
            sink.add("TFDC-605", Level.WARN,
                     f"变量全程恒定（{nums[0]!r}），疑似死点", location=loc)
        # 分布合理性（§2.3）
        if var.unit == "1":
            sorted_nums = sorted(nums)
            p95 = sorted_nums[min(len(sorted_nums) - 1,
                                  int(0.95 * (len(sorted_nums) - 1)) + 1)]
            if p95 > 1.5:
                sink.add("TFDC-606", Level.WARN,
                         f"声明为无量纲比率（1）但 p95={p95:g} > 1.5，"
                         "疑似 0–100 刻度", location=loc)
        elif var.unit == "%" and max(nums) <= 1.0:
            sink.add("TFDC-606", Level.WARN,
                     "声明为百分比（%）但最大值 ≤ 1.0，疑似 0–1 刻度",
                     location=loc)
    elif var.dtype == "boolean":
        if all(v == non_null[0] for v in non_null):
            sink.add("TFDC-605", Level.WARN,
                     f"变量全程恒定（{non_null[0]!r}），疑似死点", location=loc)


def _check_derived(
    dataset: TfdcDataset,
    registry: TfomRegistry,
    object_model: Mapping[str, str],
    columns: dict[str, list[Any]],
    options: ImportOptions,
    sink: DiagnosticSink,
) -> None:
    by_object: dict[str, list[VariableRecord]] = {}
    for var in dataset.variables:
        by_object.setdefault(var.object_id, []).append(var)
    for object_id, vars_ in by_object.items():
        model = registry.get(object_model.get(object_id, ""))
        if model is None:
            continue
        present = {v.property_code for v in vars_ if v.variable_id in columns}
        for var in vars_:
            prop = model.properties.get(var.property_code)
            if prop is None or not prop.expression:
                continue
            if var.variable_id not in columns:
                continue
            deps = expression_dependencies(prop.expression)
            if not deps <= present:
                continue  # 依赖不全，无法检查
            actual = columns[var.variable_id]
            dep_cols = {d: columns[f"{object_id}.{d}"] for d in deps}
            violations = 0
            for i, observed in enumerate(actual):
                if observed is None:
                    continue
                env = {d: dep_cols[d][i] for d in deps}
                if any(v is None for v in env.values()):
                    continue
                try:
                    expected = _eval_expression(prop.expression,
                                                {k: float(v) for k, v in env.items()})
                except (ValueError, ZeroDivisionError):
                    continue
                tol = max(options.derived_abs_tol,
                          options.derived_rel_tol * abs(expected))
                if abs(float(observed) - expected) > tol:
                    violations += 1
            if violations:
                sink.add(
                    "TFDC-604", Level.ERROR,
                    f"派生量与 TFOM expression 不一致: {prop.expression}",
                    location=f"data!{var.variable_id}", count=violations,
                )


def _build_arrow(
    timestamps: list[datetime],
    columns: dict[str, list[Any]],
    column_order: list[str],
    variables: Mapping[str, VariableRecord],
) -> pa.Table:
    arrays = [pa.array(timestamps, type=pa.timestamp("us", tz="UTC"))]
    names = ["timestamp"]
    for col in column_order:
        names.append(col)
        arrays.append(pa.array(columns[col],
                               type=_ARROW_TYPES[variables[col].dtype]))
    return pa.table(dict(zip(names, arrays)))
