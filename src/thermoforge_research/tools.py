"""Agent 工具层（architecture.md §3/§6、implementation-notes.md §11）。

architecture §6 工具清单的确定性实现，薄封装 Phase 1/2/4 能力：

- 数据：tf_dataset_import / list / get / schema / profile / query /
  sample / materialize / compare
- 研究：tf_goal_create / tf_goal_get / tf_research_status / tf_hypothesis_create /
  tf_experiment_plan / tf_experiment_run / tf_experiment_get /
  tf_model_compare / tf_model_publish
- 文献：tf_literature_search（外部学术 API 检索，只读）

约定：

- 每个工具返回统一信封（`envelope.py`）：ok/tool/id/status/inputs/
  summary/diagnostics/artifacts/truncated；有副作用的工具返回稳定 ID
  （dataset ref、RG-/H-/EXP-/VIEW-/model@version），Agent 后续引用
  ID 而非重新描述输入。
- Agent 只拿到摘要与有限样本：sample 硬上限 200 行，profile 分位数
  点位固定，query 只返回聚合统计；响应体 32 KB 上限截断。
- 已知错误（VaultError / ResearchError / ModelRegistryError / 契约
  校验错误）不抛出，转为 ok=False 的信封；诊断携带原错误码。
"""

from __future__ import annotations

import json
import math
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

import pandas as pd
from pydantic import ValidationError

from thermoforge_core.contracts.experiment import Experiment
from thermoforge_core.contracts.research_goal import ResearchGoal
from thermoforge_core.timeutil import parse_timestamp
from thermoforge_data.derive import (
    DeriveError,
    DeriveRule,
    apply_rules,
    build_dataset as build_derived_dataset,
    build_object_model,
    build_registry as build_derive_registry,
    divergence_report,
    lineage as derive_lineage,
    resolve_metadata,
    to_columns as derive_to_columns,
)
from thermoforge_data.importer import TfomRegistry, default_registry, import_parsed, import_xlsx
from thermoforge_data.legacy import convert_legacy_tables, load_workbook_model
from thermoforge_data.preprocess import (
    TFPP_APPROVAL_FORBIDDEN,
    ApprovalRecord,
    PreprocessError,
    RuleSet,
    RuleStore,
    apply_ruleset,
    validate_rule_params,
)
from thermoforge_data.vault import DataVault, VaultError
from thermoforge_data.views import (
    materialize_view,
    validate_view_definition,
    view_hash,
)
from thermoforge_research.envelope import (
    SAMPLE_MAX_ROWS,
    diagnostic_dicts,
    finalize_envelope,
    make_envelope,
)
from thermoforge_research.ledger import ResearchLedger
from thermoforge_research.runner import current_environment_lock, run_experiment
from thermoforge_research.splits import DEFAULT_EMBARGO_SECONDS
from thermoforge_runtime.errors import ModelRegistryError
from thermoforge_runtime.package import boundary_rows, build_model_package
from thermoforge_runtime.registry import ModelRegistry

from .errors import ResearchError
from . import litsearch
from . import usage
from .model_lab import (
    LabError,
    LabStore,
    parse_lab_ref,
    scan_source,
    validate_module,
)
from .modelability import ModelabilityConfig, build_modelability_report
from .whitelist import check_candidate_inputs, check_view_within_whitelist

QUERY_AGGREGATIONS = ("count", "missing", "mean", "min", "max", "std")

# tf_dataset_sample 未指定列时的默认上限（防 32KB 截断掩盖列选择失误）
SAMPLE_DEFAULT_MAX_COLUMNS = 12


class ToolContext:
    """工具调用的运行上下文（控制面持有一次，逐工具传入）。

    ::

        ctx = ToolContext(vault_root="vault", research_root="research",
                          models_root="models")
        env = tf_dataset_import(ctx, "data/workbook.xlsx")
    """

    def __init__(
        self,
        *,
        vault_root: str | Path,
        research_root: str | Path,
        models_root: str | Path | None = None,
        tfom_registry: TfomRegistry | None = None,
        actor: str = "agent",
    ):
        # 一律解析为绝对路径：实验在子进程里跑，且 cwd 会被切到实验目录，
        # 相对根会在子进程内被二次解析成 `<exp_dir>/<relative_root>/...`。
        self.vault_root = Path(vault_root).resolve()
        self.research_root = Path(research_root).resolve()
        self.models_root = (
            Path(models_root).resolve() if models_root
            else self.research_root.parent / "models"
        )
        self.tfom_registry = tfom_registry
        self.actor = actor
        self.artifacts_root = self.research_root / "tool_artifacts"
        self._vault: DataVault | None = None
        self._ledger: ResearchLedger | None = None
        self._registry: ModelRegistry | None = None
        self._preprocess_store: RuleStore | None = None
        self._lab_store: LabStore | None = None

    @property
    def vault(self) -> DataVault:
        if self._vault is None:
            self._vault = DataVault(self.vault_root)
        return self._vault

    @property
    def ledger(self) -> ResearchLedger:
        if self._ledger is None:
            self._ledger = ResearchLedger(self.research_root)
        return self._ledger

    @property
    def registry(self) -> ModelRegistry:
        if self._registry is None:
            self._registry = ModelRegistry(self.models_root)
        return self._registry

    @property
    def view_cache_root(self) -> Path:
        return self.research_root / "view_cache"

    @property
    def preprocess_store(self) -> RuleStore:
        if self._preprocess_store is None:
            self._preprocess_store = RuleStore(self.research_root / "preprocess")
        return self._preprocess_store

    @property
    def lab_store(self) -> LabStore:
        if self._lab_store is None:
            self._lab_store = LabStore(self.research_root)
        return self._lab_store


def _finish(ctx: ToolContext, env: dict[str, Any]) -> dict[str, Any]:
    return finalize_envelope(env, ctx.artifacts_root)


def _error_envelope(
    ctx: ToolContext, tool: str, exc: Exception,
    *, inputs: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """已知错误 → ok=False 信封（诊断携带原错误码，不用自由文本猜测）。"""
    code = getattr(exc, "code", None) or type(exc).__name__
    return _finish(ctx, make_envelope(
        tool, ok=False, status="FAILED", inputs=inputs,
        summary={"error": str(exc)},
        diagnostics=[{"code": code, "level": "ERROR", "count": 1,
                      "message": str(exc)[:500]}],
    ))


def _artifact(path: Path, kind: str) -> dict[str, Any]:
    return {"kind": kind, "path": str(path),
            "bytes": path.stat().st_size if path.exists() else 0}


def _rev_dir(ctx: ToolContext, ref: str) -> Path:
    return ctx.vault.resolve(ref).path


def _read_json(path: Path) -> Any:
    with open(path, encoding="utf-8") as fp:
        return json.load(fp)


# ---------------------------------------------------------------- 数据工具


def tf_dataset_import(
    ctx: ToolContext,
    path: str | Path,
    *,
    lineage: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """导入 TFDC-XLSX 并落 Data Vault，返回稳定 `dataset_id@rev_NNNN`。"""
    tool = "tf_dataset_import"
    result = import_xlsx(path, registry=ctx.tfom_registry)
    diagnostics = diagnostic_dicts(result.diagnostics)
    if not result.ok:
        return _finish(ctx, make_envelope(
            tool, ok=False, status="FAILED",
            inputs={"path": str(path)},
            summary={"error": "存在 ERROR 级诊断，未写入 vault"},
            diagnostics=diagnostics,
        ))
    try:
        ref = ctx.vault.store(result, source_path=path, lineage=lineage)
    except VaultError as exc:
        return _error_envelope(ctx, tool, exc, inputs={"path": str(path)})
    rev = _rev_dir(ctx, ref)
    has_warn = any(d["level"] in ("WARN", "REJECT") for d in diagnostics)
    table = result.table
    assert table is not None and result.dataset is not None
    ts = table.column("timestamp").to_pylist()
    return _finish(ctx, make_envelope(
        tool, ok=True, id=ref,
        status="IMPORTED_WITH_WARNINGS" if has_warn else "IMPORTED",
        inputs={"path": str(path), "importer_version": result.importer_version},
        summary={
            "dataset_ref": ref,
            "rows": table.num_rows,
            "variables": len(result.dataset.variables),
            "objects": len(result.dataset.objects),
            "time_range": [ts[0].isoformat() if ts else None,
                           ts[-1].isoformat() if ts else None],
            "degradations": list(result.degradations),
        },
        diagnostics=diagnostics,
        artifacts=[_artifact(rev / "quality.json", "quality_report"),
                   _artifact(rev / "profile.json", "profile"),
                   _artifact(rev / "manifest.json", "manifest")],
    ))


def tf_dataset_list(ctx: ToolContext) -> dict[str, Any]:
    """列出全部数据集及 revision 摘要。"""
    datasets = ctx.vault.list_datasets()
    return _finish(ctx, make_envelope(
        "tf_dataset_list",
        summary={"datasets": datasets, "count": len(datasets)},
    ))


def tf_dataset_get(ctx: ToolContext, ref: str) -> dict[str, Any]:
    """单个数据版本的清单、质量与谱系摘要。"""
    tool = "tf_dataset_get"
    try:
        rev = _rev_dir(ctx, ref)
    except VaultError as exc:
        return _error_envelope(ctx, tool, exc, inputs={"ref": ref})
    manifest = _read_json(rev / "manifest.json")
    quality = _read_json(rev / "quality.json")
    lineage = _read_json(rev / "lineage.json")
    return _finish(ctx, make_envelope(
        tool, id=ref, status="OK", inputs={"ref": ref},
        summary={
            "manifest": manifest,
            "quality": {
                "imported_at": quality.get("imported_at"),
                "degradations": quality.get("degradations", []),
                "diagnostics": quality.get("diagnostics", []),
            },
            "lineage": lineage,
        },
        diagnostics=quality.get("diagnostics", []),
        artifacts=[_artifact(rev / "quality.json", "quality_report"),
                   _artifact(rev / "profile.json", "profile")],
    ))


def tf_dataset_schema(ctx: ToolContext, ref: str) -> dict[str, Any]:
    """数据版本的变量 Schema（variable_id/unit/dtype/role/范围）。

    宽表数据集的完整 variables 明细始终落 artifact；信封 summary 携带
    紧凑清单（variable_ids/property_codes，供白名单与预检使用）+
    前 50 条明细（§11 响应上限）。
    """
    tool = "tf_dataset_schema"
    try:
        variables = ctx.vault.load_variables(ref)
        objects = ctx.vault.load_objects(ref)
    except VaultError as exc:
        return _error_envelope(ctx, tool, exc, inputs={"ref": ref})
    out_dir = ctx.artifacts_root / "schema"
    out_dir.mkdir(parents=True, exist_ok=True)
    artifact_path = out_dir / f"{ref.replace('@', '_')}-schema.json"
    fd, tmp = tempfile.mkstemp(dir=out_dir, prefix=".tmp_", suffix=".json")
    with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as fp:
        json.dump({"dataset_ref": ref, "variables": variables,
                   "objects": objects}, fp, ensure_ascii=False,
                  sort_keys=True, indent=2)
        fp.write("\n")
    os.replace(tmp, artifact_path)
    return _finish(ctx, make_envelope(
        tool, id=ref, status="OK", inputs={"ref": ref},
        summary={"variable_count": len(variables),
                 "variable_ids": sorted(v["variable_id"] for v in variables),
                 "property_codes": sorted({v["property_code"]
                                           for v in variables}),
                 "objects": objects,
                 "variables": variables[:50],
                 "variables_in_artifact": len(variables)},
        artifacts=[_artifact(artifact_path, "schema")],
    ))


def tf_dataset_profile(ctx: ToolContext, ref: str) -> dict[str, Any]:
    """数据画像摘要：分位数点位固定 min/p1/p25/p50/p75/p99/max（§11）。"""
    tool = "tf_dataset_profile"
    try:
        rev = _rev_dir(ctx, ref)
    except VaultError as exc:
        return _error_envelope(ctx, tool, exc, inputs={"ref": ref})
    profile = _read_json(rev / "profile.json")
    variables = []
    for v in profile.get("variables", []):
        entry = {
            "variable_id": v["variable_id"],
            "unit": v.get("unit"),
            "dtype": v.get("dtype"),
            "role": v.get("role"),
            "count": v.get("count"),
            "missing_rate": v.get("missing_rate"),
        }
        if "quantiles" in v:
            q = v["quantiles"]
            entry["distribution"] = {
                "min": v.get("min"),
                "p01": q.get("p01"),
                "p25": q.get("p25"),
                "p50": q.get("p50"),
                "p75": q.get("p75"),
                "p99": q.get("p99"),
                "max": v.get("max"),
            }
            entry["outlier_count_iqr"] = v.get("outlier_count_iqr")
            entry["constant"] = v.get("constant")
        variables.append(entry)
    return _finish(ctx, make_envelope(
        tool, id=ref, status="OK", inputs={"ref": ref},
        summary={
            "dataset_id": profile.get("dataset_id"),
            "record_count": profile.get("record_count"),
            "time_range": profile.get("time_range"),
            "interval_stats": profile.get("interval_stats"),
            "gap_count": profile.get("gap_count"),
            "range_violations": profile.get("range_violations"),
            "coverage": profile.get("coverage"),
            "quantile_points": ["min", "p01", "p25", "p50", "p75", "p99", "max"],
            "variables": variables,
        },
        diagnostics=profile.get("diagnostics_summary", []),
        artifacts=[_artifact(rev / "profile.json", "profile")],
    ))


def _load_frame(
    ctx: ToolContext, ref: str,
    variable_ids: Sequence[str] | None,
    start: str | None,
    end: str | None,
) -> pd.DataFrame:
    """加载 canonical 数据并按时间窗过滤（工具内部，不暴露给 Agent）。"""
    table = ctx.vault.load_data(ref)
    df = table.to_pandas()
    if start is not None:
        df = df[df["timestamp"] >= parse_timestamp(start)]
    if end is not None:
        df = df[df["timestamp"] < parse_timestamp(end)]
    if variable_ids:
        known = set(table.column_names)
        unknown = [v for v in variable_ids if v not in known]
        if unknown:
            raise VaultError("TFV-802", f"变量在该数据版本中不存在: {unknown}")
        df = df[["timestamp", *variable_ids]]
    return df


def tf_dataset_query(
    ctx: ToolContext,
    ref: str,
    *,
    variable_ids: Sequence[str],
    start: str | None = None,
    end: str | None = None,
    aggregations: Sequence[str] = ("count", "missing", "mean", "min", "max"),
) -> dict[str, Any]:
    """聚合查询：只返回统计量，不返回原始行（§6：摘要不刷屏）。"""
    tool = "tf_dataset_query"
    bad = [a for a in aggregations if a not in QUERY_AGGREGATIONS]
    if bad:
        return _finish(ctx, make_envelope(
            tool, ok=False, status="FAILED", inputs={"ref": ref},
            summary={"error": f"未登记的聚合: {bad}（允许 {QUERY_AGGREGATIONS}）"},
        ))
    try:
        df = _load_frame(ctx, ref, variable_ids, start, end)
    except (VaultError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs={"ref": ref})
    stats: dict[str, Any] = {}
    for var in variable_ids:
        col = df[var]
        non_null = col.dropna()
        entry: dict[str, Any] = {}
        for agg in aggregations:
            if agg == "count":
                entry["count"] = int(len(non_null))
            elif agg == "missing":
                entry["missing"] = int(len(col) - len(non_null))
            elif not len(non_null):
                entry[agg] = None
            elif agg == "mean":
                entry["mean"] = float(non_null.mean())
            elif agg == "min":
                entry["min"] = float(non_null.min())
            elif agg == "max":
                entry["max"] = float(non_null.max())
            elif agg == "std":
                entry["std"] = float(non_null.std()) if len(non_null) > 1 else 0.0
        stats[var] = entry
    return _finish(ctx, make_envelope(
        tool, id=ref, status="OK",
        inputs={"ref": ref, "start": start, "end": end,
                "aggregations": list(aggregations)},
        summary={"matched_rows": int(len(df)), "stats": stats},
    ))


def _json_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    return value


def tf_dataset_sample(
    ctx: ToolContext,
    ref: str,
    *,
    n: int = 50,
    variable_ids: Sequence[str] | None = None,
    start: str | None = None,
    end: str | None = None,
) -> dict[str, Any]:
    """有限样本：硬上限 200 行（§11），等距抽样（确定性）。"""
    tool = "tf_dataset_sample"
    capped = int(n) > SAMPLE_MAX_ROWS
    take = min(max(int(n), 1), SAMPLE_MAX_ROWS)
    try:
        df = _load_frame(ctx, ref, variable_ids, start, end)
    except (VaultError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs={"ref": ref})
    if variable_ids is None and len(df.columns) > SAMPLE_DEFAULT_MAX_COLUMNS + 1:
        df = df.iloc[:, : SAMPLE_DEFAULT_MAX_COLUMNS + 1]
    total = len(df)
    if total > take:
        idx = [round(i * (total - 1) / (take - 1)) for i in range(take)] \
            if take > 1 else [0]
        df = df.iloc[sorted(set(idx))]
    rows = [
        {k: _json_value(v) for k, v in row.items()}
        for row in df.to_dict(orient="records")
    ]
    env = make_envelope(
        tool, ok=True, id=ref, status="OK",
        inputs={"ref": ref, "requested": int(n), "start": start, "end": end},
        summary={
            "rows": rows,
            "returned": len(rows),
            "total_matched": int(total),
            "cap": SAMPLE_MAX_ROWS,
        },
    )
    if capped:
        env["truncated"] = True  # 策略性截断：请求超硬上限
    return _finish(ctx, env)


def tf_dataset_materialize(
    ctx: ToolContext,
    view_definition: Mapping[str, Any],
) -> dict[str, Any]:
    """登记并物化 Dataset View，返回稳定 VIEW-ID 与 view_hash。"""
    tool = "tf_dataset_materialize"
    doc = dict(view_definition)
    try:
        validate_view_definition(doc)
        mv = materialize_view(ctx.vault, doc, ctx.view_cache_root)
        view = ctx.ledger.register_view(
            {**doc, "view_hash": mv.view_hash}, actor=ctx.actor,
            reason="工具层登记 Dataset View",
        )
    except (VaultError, ValueError, KeyError) as exc:
        return _error_envelope(ctx, tool, exc,
                               inputs={"dataset": doc.get("dataset")})
    return _finish(ctx, make_envelope(
        tool, id=view["id"], status="MATERIALIZED",
        inputs={"dataset": doc.get("dataset"), "view_hash": mv.view_hash},
        summary={
            "view_id": view["id"],
            "view_hash": mv.view_hash,
            "rows": mv.table.num_rows,
            "columns": mv.table.column_names,
            "cache_reused": mv.reused,
        },
        artifacts=[_artifact(mv.path / "data.parquet", "view_data"),
                   _artifact(mv.path / "view.json", "view_definition")],
    ))


def _derived_model_id(object_id: str) -> str:
    """object_id → 合法的 object_model_id（`^[a-z][a-z0-9_]*\\.v[0-9]+$`）。"""
    slug = "".join(c if c.isalnum() else "_" for c in object_id.lower())
    slug = slug.strip("_") or "derived"
    if not slug[0].isalpha():
        slug = f"d_{slug}"
    return f"{slug}.v1"


def tf_dataset_derive(
    ctx: ToolContext,
    source_ref: str,
    *,
    dataset_id: str,
    object_id: str,
    rules: Sequence[Mapping[str, Any]],
    object_model: str | None = None,
    object_name: str | None = None,
    description: str | None = None,
) -> dict[str, Any]:
    """按声明式规则派生新数据集并落 vault，返回稳定 `dataset_id@rev_NNNN`。

    跨对象聚合（「两路总管流量之和」「运行设备功率之和」）既不能用单机
    物模型表达，也不在 Dataset View 的长表机制内（implementation-notes
    §3.3），此前只能写一次性脚本。本工具把它变成可由 Agent 调用的参数化
    能力：算子取自 `thermoforge_data.derive.OPS` 白名单，**不接受自由
    代码**；产出走标准导入管线并记录完整派生谱系（I-49）。
    """
    tool = "tf_dataset_derive"
    inputs = {"source_ref": source_ref, "dataset_id": dataset_id,
              "object_id": object_id, "rule_count": len(rules)}
    try:
        parsed = [DeriveRule.parse(r) for r in rules]
        source_vars = ctx.vault.load_variables(source_ref)
        units, dtypes = resolve_metadata(parsed, source_vars)
        df = ctx.vault.load_data(source_ref).to_pandas()
        frame = apply_rules(df, parsed)
        divergence = divergence_report(df, parsed)
        model_id = object_model or _derived_model_id(object_id)
        model = build_object_model(model_id, parsed, units, dtypes)
        dataset = build_derived_dataset(
            dataset_id, object_id, model_id, parsed, units, dtypes,
            _read_json(_rev_dir(ctx, source_ref) / "manifest.json") or {},
            object_name, description)
        result = import_parsed(
            dataset, frame["timestamp"].tolist(),
            derive_to_columns(frame, object_id, dtypes),
            registry=build_derive_registry(model, ctx.tfom_registry),
            degradations=[f"派生自 {source_ref}（{len(parsed)} 条规则）"])
    except (DeriveError, VaultError, ValueError, KeyError) as exc:
        return _error_envelope(ctx, tool, exc, inputs=inputs)

    diagnostics = diagnostic_dicts(result.diagnostics)
    if not result.ok:
        return _finish(ctx, make_envelope(
            tool, ok=False, status="FAILED", inputs=inputs,
            summary={"error": "派生结果未通过 TFDC 校验，未写入 vault"},
            diagnostics=diagnostics))
    try:
        ref = ctx.vault.store(
            result, lineage=derive_lineage(source_ref, parsed, divergence))
    except VaultError as exc:
        return _error_envelope(ctx, tool, exc, inputs=inputs)

    table = result.table
    assert table is not None
    non_null = {
        rule.property_code: int(
            table.column(f"{object_id}.{rule.property_code}").length()
            - table.column(f"{object_id}.{rule.property_code}").null_count)
        for rule in parsed
    }
    return _finish(ctx, make_envelope(
        tool, ok=True, id=ref,
        status="DERIVED_WITH_WARNINGS" if diagnostics else "DERIVED",
        inputs=inputs,
        summary={
            "rows": table.num_rows,
            "properties": {r.property_code: r.spec_text() for r in parsed},
            "units": dict(units),
            "non_null_counts": non_null,
            "redundancy_check": divergence,
            "derived_from": source_ref,
        },
        diagnostics=diagnostics))


def tf_dataset_compare(ctx: ToolContext, ref_a: str, ref_b: str) -> dict[str, Any]:
    """两个数据版本的对比：行数/时间范围/变量增删/数值漂移（有界）。"""
    tool = "tf_dataset_compare"
    try:
        info_a, info_b = ctx.vault.resolve(ref_a), ctx.vault.resolve(ref_b)
        vars_a = {v["variable_id"]: v for v in ctx.vault.load_variables(ref_a)}
        vars_b = {v["variable_id"]: v for v in ctx.vault.load_variables(ref_b)}
    except VaultError as exc:
        return _error_envelope(ctx, tool, exc,
                               inputs={"ref_a": ref_a, "ref_b": ref_b})
    df_a = ctx.vault.load_data(ref_a).to_pandas()
    df_b = ctx.vault.load_data(ref_b).to_pandas()
    common = sorted(set(vars_a) & set(vars_b))
    drift = []
    for var in common:
        if vars_a[var].get("dtype") not in ("float", "integer"):
            continue
        a = pd.to_numeric(df_a[var], errors="coerce").dropna()
        b = pd.to_numeric(df_b[var], errors="coerce").dropna()
        if not len(a) or not len(b):
            continue
        mean_a, mean_b = float(a.mean()), float(b.mean())
        drift.append({
            "variable_id": var, "mean_a": mean_a, "mean_b": mean_b,
            "rel_change": (abs(mean_b - mean_a) / max(abs(mean_a), 1e-12)),
        })
    drift.sort(key=lambda d: -d["rel_change"])
    return _finish(ctx, make_envelope(
        tool, status="OK", inputs={"ref_a": ref_a, "ref_b": ref_b},
        summary={
            "content_identical": info_a.content_sha256 == info_b.content_sha256,
            "rows": [int(len(df_a)), int(len(df_b))],
            "time_range_a": [str(df_a["timestamp"].min()),
                             str(df_a["timestamp"].max())],
            "time_range_b": [str(df_b["timestamp"].min()),
                             str(df_b["timestamp"].max())],
            "variables_added": sorted(set(vars_b) - set(vars_a)),
            "variables_removed": sorted(set(vars_a) - set(vars_b)),
            "variables_common": len(common),
            "drift_top": drift[:20],
        },
    ))


# ---------------------------------------------------------------- 研究工具


def tf_dataset_modelability(
    ctx: ToolContext,
    ref: str,
    *,
    goal_id: str | None = None,
    target: str | None = None,
    candidate_inputs: Sequence[str] | None = None,
    object_model: str | None = None,
    same_origin_corr: float | None = None,
) -> dict[str, Any]:
    """可建模性报告（G3 语义层检查）。FAIL = 存在 blocker，不得进入建模。

    target/candidate_inputs 可由 `goal_id` 从账本中的 Research Goal 定义
    读取；完整报告始终落 artifact，信封只带摘要（§11）。

    `same_origin_corr` 覆盖同源阻断阈值（默认 0.98，implementation-notes §14
    写明是待标定的参考值）。**只在能说清「高相关是物理耦合而非同源」时才动**：
    该检查同时兜住了没有派生元数据的循环论证（F1 `power~current_percent`
    r=0.9955），调高即放松那道网。用了什么阈值会写进报告 artifact 的
    `config` 段，事后可查。
    """
    tool = "tf_dataset_modelability"
    try:
        if goal_id is not None:
            goal = ctx.ledger.get(goal_id)
            definition = goal.get("definition") or {}
            target = target or definition.get("target")
            candidate_inputs = candidate_inputs or definition.get(
                "candidate_inputs")
            object_model = object_model or definition.get("object_model")
        if not target or not candidate_inputs:
            raise ValueError("必须提供 goal_id 或显式 target + candidate_inputs")
        config = (ModelabilityConfig(same_origin_corr=float(same_origin_corr))
                  if same_origin_corr is not None else None)
        report = build_modelability_report(
            ctx.vault, ref, target=str(target),
            candidate_inputs=[str(c) for c in candidate_inputs],
            object_model=object_model, registry=ctx.tfom_registry,
            config=config,
        )
    except (VaultError, KeyError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs={"ref": ref,
                                                       "goal_id": goal_id})
    out_dir = ctx.artifacts_root / "modelability"
    out_dir.mkdir(parents=True, exist_ok=True)
    artifact_path = out_dir / f"{ref.replace('@', '_').replace('/', '_')}-modelability.json"
    fd, tmp = tempfile.mkstemp(dir=out_dir, prefix=".tmp_", suffix=".json")
    with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as fp:
        json.dump(report, fp, ensure_ascii=False, sort_keys=True, indent=2)
        fp.write("\n")
    os.replace(tmp, artifact_path)
    ok = report["verdict"] == "PASS"
    return _finish(ctx, make_envelope(
        tool, ok=ok, id=ref, status=report["verdict"],
        inputs={"ref": ref, "goal_id": goal_id, "target": report["target"]},
        summary={
            "verdict": report["verdict"],
            "blockers": report["blockers"],
            "warnings": report["warnings"],
            "checks": [
                {"name": c["name"], "level": c["level"],
                 "passed": c["passed"], "summary": c["summary"]}
                for c in report["checks"]
            ],
        },
        artifacts=[_artifact(artifact_path, "modelability_report")],
    ))


def tf_goal_create(
    ctx: ToolContext,
    definition: Mapping[str, Any],
    *,
    dataset_ref: str | None = None,
) -> dict[str, Any]:
    """创建 Research Goal（契约校验 + 稳定 RG-ID）。

    `dataset_ref` 提供时做白名单机器校验（DD-16）：candidate_inputs
    必须存在于数据版本且不得为派生量（data-survey §F1 循环论证）。
    """
    tool = "tf_goal_create"
    doc = dict(definition)
    doc.setdefault("goal_id", ctx.ledger.allocator.allocate("RG-"))
    try:
        goal = ResearchGoal(**doc)
        if dataset_ref is not None:
            violations = check_candidate_inputs(
                goal.candidate_inputs, target=goal.target,
                variables=ctx.vault.load_variables(dataset_ref),
            )
            if violations:
                raise ValueError(
                    "candidate_inputs 白名单校验失败: " + "; ".join(violations)
                )
        entity = ctx.ledger.create_goal(
            goal.name, actor=ctx.actor,
            definition=goal.model_dump(mode="json"),
            entity_id=goal.goal_id, reason="工具层创建 Research Goal",
        )
    except (ValidationError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs={"name": doc.get("name")})
    return _finish(ctx, make_envelope(
        tool, id=entity["id"], status="CREATED",
        inputs={"goal_id": entity["id"], "dataset_ref": dataset_ref},
        summary={
            "goal_id": entity["id"],
            "name": goal.name,
            "target": goal.target,
            "object_model": goal.object_model,
            "candidate_inputs": list(goal.candidate_inputs),
            "acceptance": goal.acceptance.model_dump(mode="json",
                                                     exclude_none=True),
            "budgets": {
                "max_experiments": goal.max_experiments,
                "max_duration_days": goal.max_duration_days,
                "compute_budget_hours": goal.compute_budget_hours,
            },
            "approval_required": list(goal.approval_required),
        },
    ))


def tf_goal_get(ctx: ToolContext, goal_id: str) -> dict[str, Any]:
    """读取 Research Goal 的完整定义（**含 candidate_inputs 白名单原文**）。

    没有这个工具，Agent 就无法知道 DD-16 白名单里到底有哪些条目 —— 只能靠猜，
    而猜错会在 `tf_experiment_plan` 阶段被机器拒绝，且拒绝信息只说"不在白名单内"，
    不说白名单是什么。这是一条真实踩出来的死路（agent 会话
    `research/agent_sessions/20260814T014407*.jsonl`）。

    白名单条目有两种写法，语义不同：

    - 裸 `property_code`（如 `cooling_load`）：范围内任意对象的该属性都允许
    - `object_id.property_code`（如 `CH01.cooling_load`）：仅该对象的该属性允许，
      **且要求实验视图的对象范围包含该对象**

    实验的 `features` 一律填 `property_code`。若白名单用了带对象的写法而视图是
    多对象长表，校验会失败 —— 此时应改用裸写法，或把视图过滤到该对象。
    """
    tool = "tf_goal_get"
    try:
        entity = ctx.ledger.get(goal_id)
    except (KeyError, ResearchError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs={"goal_id": goal_id})
    if entity is None:
        return _finish(ctx, make_envelope(
            tool, ok=False, status="NOT_FOUND", inputs={"goal_id": goal_id},
            summary={"error": f"目标不存在: {goal_id}"},
        ))
    definition = dict(entity.get("definition") or {})
    entries = [str(e) for e in (definition.get("candidate_inputs") or [])]
    scoped = sorted({e.split(".", 1)[0] for e in entries if "." in e})
    return _finish(ctx, make_envelope(
        tool, id=goal_id, status="OK", inputs={"goal_id": goal_id},
        summary={
            "goal_id": goal_id,
            "name": definition.get("name") or entity.get("name"),
            "status": entity.get("status"),
            "target": definition.get("target"),
            "object_model": definition.get("object_model"),
            "purpose": definition.get("purpose"),
            "description": definition.get("description"),
            "candidate_inputs": entries,
            "whitelist_style": ("object_scoped" if scoped
                                else "bare_property" if entries else "empty"),
            "whitelist_objects": scoped,
            "acceptance": definition.get("acceptance"),
            "model_types": definition.get("model_types"),
            "approval_required": definition.get("approval_required") or [],
            "budgets": {
                k: definition.get(k) for k in
                ("max_experiments", "max_duration_days", "compute_budget_hours")
            },
        },
    ))


def tf_research_status(
    ctx: ToolContext,
    goal_id: str | None = None,
) -> dict[str, Any]:
    """研究进展：目标状态、假设/实验/模型计数、预算消耗。"""
    tool = "tf_research_status"
    if goal_id is None:
        goals = ctx.ledger.list_goals()
        return _finish(ctx, make_envelope(
            tool, status="OK",
            summary={"goals": [
                {"goal_id": g["id"], "name": g.get("name"),
                 "status": g["status"]}
                for g in goals
            ]},
        ))
    try:
        goal = ctx.ledger.get(goal_id)
        progress = ctx.ledger.goal_progress(goal_id)
    except (KeyError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs={"goal_id": goal_id})
    definition = goal.get("definition") or {}
    spent = progress["experiments"]["total"]
    budget: dict[str, Any] = {
        "experiments_used": spent,
        "experiments_max": definition.get("max_experiments"),
    }
    if definition.get("max_experiments") is not None:
        budget["experiments_remaining"] = max(
            0, int(definition["max_experiments"]) - spent)
    created = definition.get("created_at") or goal.get("created_at")
    if created and definition.get("max_duration_days"):
        elapsed = (datetime.now(timezone.utc)
                   - parse_timestamp(created)).total_seconds() / 86400.0
        budget["days_elapsed"] = elapsed
        budget["days_max"] = definition["max_duration_days"]
    return _finish(ctx, make_envelope(
        tool, id=goal_id, status=str(goal["status"]).upper(),
        inputs={"goal_id": goal_id},
        summary={
            "progress": progress,
            "budget": budget,
            "unverified_hypotheses": [
                h["id"] for h in ctx.ledger.unverified_hypotheses(goal_id)],
            "failed_experiments": ctx.ledger.failed_experiments(goal_id),
        },
    ))


# ---------------------------------------------------------------- 用量归账查询
#
# 与网页控制台「AI 研究 → 用量与留痕」页签同一份聚合（`usage.py`），
# 外部 Agent 经 MCP 也能查到和页面一致的数字与思维链。


def _usage_book(ctx: ToolContext
                ) -> tuple[usage.UsageBook, list[usage.TurnUsage]]:
    """全量归账账本 + 会话轮次。会话文件多，每次调用全量扫——这是低频
    查询工具，正确性优先；界面侧的高频读取另有 st.cache_data 增量缓存。"""
    session_dir = ctx.research_root / "agent_sessions"
    turns: list[usage.TurnUsage] = []
    scanned = 0
    if session_dir.is_dir():
        for path in sorted(session_dir.glob("*.jsonl")):
            turns.extend(usage.parse_session_file(path))
            scanned += 1
    book = usage.build_book(ctx.research_root, turns)
    book.scanned_sessions = scanned
    return book, turns


def _usage_entry_dict(entry: usage.UsageEntry, *, include_raw: bool = False
                      ) -> dict[str, Any]:
    """UsageEntry → 信封里的结构化条目（思维链带标签，raw_reply 按需）。"""
    out: dict[str, Any] = {
        "at": entry.at,
        "source": entry.source,
        "detail": entry.detail or None,
        "question": entry.question or None,
        "prompt_tokens": round(entry.prompt_tokens),
        "completion_tokens": round(entry.completion_tokens),
        "cost": entry.cost,
        "calls": round(entry.calls, 2),
        "share": entry.share,
    }
    if entry.reasoning:
        out["reasoning"] = [
            {"label": (entry.reasoning_labels[i]
                       if i < len(entry.reasoning_labels) else "思维链"),
             "text": text}
            for i, text in enumerate(entry.reasoning)
        ]
    if include_raw and entry.raw_reply:
        out["raw_reply"] = entry.raw_reply
    return out


def _usage_row(view: usage.EntityUsage) -> dict[str, Any]:
    return {
        "id": view.entity_id,
        "prompt_tokens": round(view.prompt_tokens),
        "completion_tokens": round(view.completion_tokens),
        "cost": view.cost,  # None = 未在 agent.toml 配置 [pricing] 单价
        "calls": round(view.calls, 1),
    }


_ATTRIBUTION_NOTE = (
    "规划留痕精确归到实验并沿 report.json 的 refs 上卷假设/目标；"
    "会话轮次在对对象做创建/运行/发布时均摊归账，只读轮次围绕唯一对象或"
    "同一目标时归账，跨目标总览/对比进未归账；外部 Agent（MCP）自己进程的"
    "对话 token 不在账内——账里只有经过本系统自带 Agent 的调用。")


def tf_usage_overview(ctx: ToolContext) -> dict[str, Any]:
    """Token 用量总览：按研究目标的归账汇总 + 规划/会话/未归账三块总量。

    与网页控制台「AI 研究 → 用量与留痕」页签同源同数。要某个目标或实验的
    逐条思维链，用 tf_goal_usage / tf_experiment_usage。
    """
    tool = "tf_usage_overview"
    book, turns = _usage_book(ctx)
    goal_ids = sorted(
        {g for g, _ in book.experiment_refs.values() if g}
        | {g for g in book.hypothesis_refs.values() if g}
        | {e for e in book.per_entity if e.startswith("RG-")})
    rows = []
    for goal_id in goal_ids:
        view = usage.goal_view(book, goal_id)
        if not view.total.entries:
            continue
        try:
            name = (ctx.ledger.get(goal_id) or {}).get("name")
        except (KeyError, ValueError):
            name = None
        rows.append({**_usage_row(view.total), "goal_id": goal_id,
                     "name": name,
                     "hypotheses": len(view.hypotheses),
                     "experiments": len(view.experiments)})
    planner_entries = [e for entries in book.per_entity.values()
                       for e in entries if e.source == "规划"]
    unassigned = usage.unassigned_view(book)
    return _finish(ctx, make_envelope(
        tool, status="OK",
        summary={
            "attribution": _ATTRIBUTION_NOTE,
            "sessions_scanned": book.scanned_sessions,
            "planner_traces": book.planner_traces,
            "orphan_planner_traces": book.orphan_traces,
            "planner_tokens": {
                "prompt": round(sum(e.prompt_tokens for e in planner_entries)),
                "completion": round(sum(e.completion_tokens
                                        for e in planner_entries)),
            },
            "session_tokens": {
                "prompt": round(sum(t.prompt_tokens for t in turns)),
                "completion": round(sum(t.completion_tokens for t in turns)),
                "turns": len(turns),
            },
            "unassigned": {**_usage_row(unassigned)},
            "goals": rows,
        },
    ))


def tf_goal_usage(ctx: ToolContext, goal_id: str) -> dict[str, Any]:
    """单个研究目标的归账明细：总量、按假设/实验分解、逐条思维链留痕。"""
    tool = "tf_goal_usage"
    try:
        goal = ctx.ledger.get(goal_id)
    except (KeyError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs={"goal_id": goal_id})
    book, _turns = _usage_book(ctx)
    view = usage.goal_view(book, goal_id)
    entries = sorted(view.total.entries, key=lambda e: e.at, reverse=True)
    return _finish(ctx, make_envelope(
        tool, id=goal_id, status="OK", inputs={"goal_id": goal_id},
        summary={
            "goal": {"goal_id": goal_id, "name": goal.get("name"),
                     "status": goal.get("status")},
            "attribution": _ATTRIBUTION_NOTE,
            "total": _usage_row(view.total),
            "hypotheses": [
                {**_usage_row(v),
                 "statement": (book.hypothesis_statement.get(v.entity_id)
                               or "")[:200]}
                for v in view.hypotheses],
            "experiments": [_usage_row(v) for v in view.experiments],
            # 新的在前；raw_reply 在目标级不带（太长），到 tf_experiment_usage 拿
            "entries": [_usage_entry_dict(e) for e in entries],
        },
    ))


def tf_experiment_usage(ctx: ToolContext, experiment_id: str) -> dict[str, Any]:
    """单个实验的归账明细：规划留痕（思维链/原始回复/用量）+ 会话份额。"""
    tool = "tf_experiment_usage"
    exp_dir = ctx.research_root / "experiments" / experiment_id
    if not exp_dir.is_dir():
        return _error_envelope(ctx, tool,
                               KeyError(f"实验不存在: {experiment_id}"),
                               inputs={"experiment_id": experiment_id})
    book, _turns = _usage_book(ctx)
    view = usage.experiment_view(book, experiment_id)
    goal_id, hypothesis_id = book.experiment_refs.get(
        experiment_id, (None, None))
    return _finish(ctx, make_envelope(
        tool, id=experiment_id, status="OK",
        inputs={"experiment_id": experiment_id},
        summary={
            "goal_id": goal_id,
            "hypothesis_id": hypothesis_id,
            "attribution": _ATTRIBUTION_NOTE,
            "total": _usage_row(view),
            "entries": [_usage_entry_dict(e, include_raw=True)
                        for e in view.entries],
        },
    ))


def tf_hypothesis_create(
    ctx: ToolContext,
    goal_id: str,
    statement: str,
    *,
    basis: Sequence[str] | None = None,
) -> dict[str, Any]:
    """创建假设；非首个假设必须引用已有证据（research-loop §3）。

    `basis=None` 与不传等价（首个假设可无证据）。显式传 None 曾直接抛
    TypeError —— 工具层不得抛异常，一切失败都走 ok=False 信封。
    """
    tool = "tf_hypothesis_create"
    basis = list(basis or ())
    try:
        hyp = ctx.ledger.create_hypothesis(
            goal_id, statement, actor=ctx.actor, basis=basis,
            reason="工具层创建假设",
        )
    except ValueError as exc:
        return _error_envelope(ctx, tool, exc,
                               inputs={"goal_id": goal_id, "basis": basis})
    return _finish(ctx, make_envelope(
        tool, id=hyp["id"], status="UNVERIFIED",
        inputs={"goal_id": goal_id, "basis": basis},
        summary={"hypothesis_id": hyp["id"], "statement": statement,
                 "basis": hyp["refs"].get("basis", [])},
    ))


def _fill_environment_lock(doc: dict[str, Any]) -> None:
    """`runtime.environment_lock` 缺省时填当前机器的真指纹。

    环境指纹是机器算出来的（`current_environment_lock()`），调用方无从得知。
    以前它是必填，于是 Agent 只能编一个（实测编出 `sim-2026-08-default`），
    每次 run 都撞 TFX-901 —— 一个本该防「跨机复现」的门禁，退化成了对
    「猜不中哈希」的惩罚。留空即声明「就在本机跑」，是唯一诚实的默认值。

    刻意不补 `random_seed`：种子是**研究决策**，必须由调用方明确写下，
    缺了就该报 TFX-902。两者看着都在 `runtime` 里，性质完全不同。
    """
    runtime = doc.get("runtime")
    if not isinstance(runtime, Mapping):
        return
    if str(runtime.get("environment_lock") or "").strip():
        return
    doc["runtime"] = {**runtime,
                      "environment_lock": current_environment_lock()[0]}


def tf_experiment_plan(
    ctx: ToolContext,
    definition: Mapping[str, Any],
) -> dict[str, Any]:
    """登记实验定义（契约校验 + 稳定 EXP-ID，不执行）。"""
    tool = "tf_experiment_plan"
    doc = dict(definition)
    doc.setdefault("experiment_id", ctx.ledger.allocator.allocate("EXP-"))
    _fill_environment_lock(doc)
    try:
        exp = Experiment(**doc)
        # DD-16 封闭白名单的机器校验：View 特征必须 ⊆ candidate_inputs，
        # 目标必须与 goal.target 一致（goal/view 均在账本时可查）
        goal_entity = ctx.ledger.get(exp.goal_id)
        view_entity = ctx.ledger.get(exp.dataset_view)
        candidate = (goal_entity.get("definition") or {}).get("candidate_inputs")
        goal_target = (goal_entity.get("definition") or {}).get("target")
        view_def = view_entity.get("definition") or {}
        if candidate and view_def:
            violations = check_view_within_whitelist(
                features=[str(f) for f in view_def.get("features", [])],
                view_target=(str(view_def["target"])
                             if view_def.get("target") else None),
                view_objects=[str(o) for o in view_def.get("objects", [])],
                candidate_inputs=[str(c) for c in candidate],
                goal_target=str(goal_target),
            )
            if violations:
                raise ValueError("DD-16 白名单校验失败: " + "; ".join(violations))
        entity = ctx.ledger.register_experiment(
            exp.model_dump(by_alias=True, mode="json"), actor=ctx.actor,
            reason="工具层登记实验计划",
        )
    except (ValidationError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc,
                               inputs={"goal_id": doc.get("goal_id")})
    return _finish(ctx, make_envelope(
        tool, id=entity["id"], status="PLANNED",
        inputs={"goal_id": exp.goal_id, "hypothesis_id": exp.hypothesis_id,
                "dataset_view": exp.dataset_view},
        summary={
            "experiment_id": entity["id"],
            "model": exp.model.model_dump(mode="json"),
            "target": exp.target,
            "random_seed": exp.runtime.random_seed,
            # 回报实际登记的指纹：填缺省时调用方要看得见填了什么
            "environment_lock": exp.runtime.environment_lock,
        },
    ))


def _experiment_artifacts(ctx: ToolContext, exp_id: str) -> list[dict[str, Any]]:
    exp_dir = ctx.research_root / "experiments" / exp_id
    out = []
    for name, kind in (("report.json", "experiment_report"),
                       ("metrics.json", "metrics"),
                       ("predictions.parquet", "predictions"),
                       ("split.json", "split"),
                       ("split_profile.json", "split_profile"),
                       ("physics_report.json", "physics_report"),
                       ("planner_trace.json", "planner_trace")):
        path = exp_dir / name
        if path.exists():
            out.append(_artifact(path, kind))
    if (exp_dir / "model").is_dir():
        out.append({"kind": "model_dir", "path": str(exp_dir / "model"),
                    "bytes": 0})
    return out


def _metrics_summary(report: Mapping[str, Any]) -> dict[str, Any]:
    metrics = report.get("metrics") or {}
    surfaces = {
        name: {"n_samples": surf.get("n_samples"),
               "metrics": surf.get("metrics")}
        for name, surf in (metrics.get("surfaces") or {}).items()
    }
    physics = report.get("physics") or {}
    # 滚动交叉验证的折间汇总必须进信封：它是一整类 Goal 声明的**判据面**
    # （acceptance.evaluated_on=rolling_cv）。信封里没有，编排器与 Agent
    # 就只能拿面 A/validate 去判 —— 判据面声明了却判在别的面上，等于没声明
    # （实测 EXP-0090：滚动 R²=0.978 达标，却因信封里读不到而没触发验收）。
    # 只带汇总数值，逐折明细仍在 metrics.json artifact 里。
    rolling = metrics.get("rolling_cv") or {}
    return {
        "surfaces": surfaces,
        "rolling_cv": ({"n_folds": rolling.get("n_folds"),
                        "n_folds_skipped": rolling.get("n_folds_skipped"),
                        "n_samples": rolling.get("n_samples"),
                        "metrics": rolling.get("metrics")}
                       if rolling.get("n_folds") else None),
        "physics_overall_rate": physics.get("overall_rate"),
        "duration_seconds": report.get("duration_seconds"),
        "environment_lock": report.get("environment_lock"),
    }


def tf_experiment_run(
    ctx: ToolContext,
    experiment_id: str,
    *,
    purge_seconds: float = 0.0,
    embargo_seconds: float = DEFAULT_EMBARGO_SECONDS,
    y_floor: float | None = None,
    timeout_seconds: float | None = None,
) -> dict[str, Any]:
    """执行已登记的实验（子进程隔离），返回结构化指标摘要。"""
    tool = "tf_experiment_run"
    try:
        entity = ctx.ledger.get(experiment_id)
        exp = Experiment(**entity["definition"])
        report = run_experiment(
            exp, research_root=ctx.research_root, vault_root=ctx.vault_root,
            ledger=ctx.ledger, purge_seconds=purge_seconds,
            embargo_seconds=embargo_seconds, y_floor=y_floor,
            actor=ctx.actor, timeout_seconds=timeout_seconds,
        )
    except (KeyError, ValidationError, ValueError, ResearchError) as exc:
        return _error_envelope(ctx, tool, exc,
                               inputs={"experiment_id": experiment_id})
    ok = report["status"] == "completed"
    diagnostics = []
    if not ok:
        diagnostics.append({
            "code": report.get("error_code") or "CHILD_ERROR",
            "level": "ERROR", "count": 1,
            "message": str(report.get("error"))[:500],
        })
    return _finish(ctx, make_envelope(
        tool, ok=ok, id=experiment_id,
        status=report["status"].upper(),
        inputs={"experiment_id": experiment_id,
                "random_seed": report.get("random_seed")},
        summary={
            **_metrics_summary(report),
            "error_code": report.get("error_code"),
            "conclusion": report.get("conclusion"),
        },
        diagnostics=diagnostics,
        artifacts=_experiment_artifacts(ctx, experiment_id),
    ))


def _fitted_parameters(exp_dir: Path) -> dict[str, Any] | None:
    """读出物理模型辨识出的参数（`model/model.json` 的 `parameters`）。

    必须进 summary：光有指标无法判断辨识是否成功 —— 参数贴死在边界
    （该项不可辨识）或量级失控（模型被推到饱和区）都不会让 CVRMSE 变差，
    只看指标会把这两种病症放过去。缺了它，「参数体检」这一步做不了，
    Agent 只能如实说"读不到"（实测 EXP-0022 会话）。
    """
    path = exp_dir / "model" / "model.json"
    if not path.is_file():
        return None
    try:
        doc = _read_json(path)
    except (OSError, json.JSONDecodeError):
        return None
    params = doc.get("parameters")
    if not isinstance(params, Mapping) or not params:
        return None
    return {"format": doc.get("format"), "values": dict(params),
            "n_train": doc.get("n_train")}


def tf_experiment_get(ctx: ToolContext, experiment_id: str) -> dict[str, Any]:
    """实验结果查询（Ledger 状态 + report.json 摘要 + 已辨识的物理参数）。"""
    tool = "tf_experiment_get"
    try:
        entity = ctx.ledger.get(experiment_id)
    except (KeyError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc,
                               inputs={"experiment_id": experiment_id})
    exp_dir = ctx.research_root / "experiments" / experiment_id
    report_path = exp_dir / "report.json"
    summary: dict[str, Any] = {
        "experiment_id": experiment_id,
        "status": entity["status"],
        "refs": entity.get("refs", {}),
    }
    artifacts: list[dict[str, Any]] = []
    if report_path.exists():
        report = _read_json(report_path)
        summary.update(_metrics_summary(report))
        summary["error_code"] = report.get("error_code")
        summary["conclusion"] = report.get("conclusion")
        params = _fitted_parameters(exp_dir)
        if params is not None:
            summary["fitted_parameters"] = params
        artifacts = _experiment_artifacts(ctx, experiment_id)
    return _finish(ctx, make_envelope(
        tool, ok=entity["status"] != "failed", id=experiment_id,
        status=str(entity["status"]).upper(),
        inputs={"experiment_id": experiment_id},
        summary=summary, artifacts=artifacts,
    ))


def _primary_surface_name(metrics: Mapping[str, Any],
                          evaluated_on: str = "auto") -> str | None:
    """发布/比较口径：默认 C→A→validate（§4.3），可由 Goal 钉死。

    必须与 `ModelRegistry._primary_surface` 同调 —— 比较排出来的第一名
    要是发布门禁判的不是同一个面，「最好的模型」就成了两件事。
    """
    metrics = metrics or {}
    if evaluated_on == "rolling_cv":
        return "rolling_cv" if (metrics.get("rolling_cv") or {}).get(
            "n_folds") else None
    surfaces = metrics.get("surfaces") or {}
    order = (("C", "A", "validate") if evaluated_on in ("auto", "", None)
             else (evaluated_on,))
    for name in order:
        if (surfaces.get(name) or {}).get("n_samples"):
            return name
    return None


def _goal_evaluated_on(ctx: ToolContext, goal_id: Any) -> str:
    """Goal 声明的判据面；Goal 读不到时退回 auto（不因此让比较失败）。"""
    if not goal_id:
        return "auto"
    try:
        goal = ctx.ledger.get(str(goal_id))
    except (KeyError, ValueError):
        return "auto"
    acceptance = (goal.get("definition") or {}).get("acceptance") or {}
    return str(acceptance.get("evaluated_on") or "auto")


def _surface_metrics(metrics: Mapping[str, Any],
                     surface: str | None) -> dict[str, Any]:
    if not surface:
        return {}
    if surface == "rolling_cv":
        return dict((metrics.get("rolling_cv") or {}).get("metrics") or {})
    return dict(((metrics.get("surfaces") or {}).get(surface) or {})
                .get("metrics") or {})


def tf_model_compare(
    ctx: ToolContext,
    experiment_ids: Sequence[str],
    *,
    evaluated_on: str | None = None,
) -> dict[str, Any]:
    """模型比较：按主测试面 CVRMSE 排名，附物理违规率。

    判据面默认取各实验所属 Goal 的 `acceptance.evaluated_on`（缺省 auto =
    C→A→validate）；`evaluated_on` 参数可临时覆盖，用于「换个口径看看」。
    """
    tool = "tf_model_compare"
    rows: list[dict[str, Any]] = []
    skipped: list[str] = []
    for exp_id in experiment_ids:
        report_path = (ctx.research_root / "experiments" / str(exp_id)
                       / "report.json")
        if not report_path.exists():
            skipped.append(str(exp_id))
            continue
        report = _read_json(report_path)
        if report.get("status") != "completed":
            skipped.append(str(exp_id))
            continue
        criterion = evaluated_on or _goal_evaluated_on(ctx, report.get("goal_id"))
        metrics = report.get("metrics") or {}
        surface = _primary_surface_name(metrics, criterion)
        rows.append({
            "experiment_id": str(exp_id),
            "primary_surface": surface,
            "metrics": _surface_metrics(metrics, surface),
            "physics_overall_rate": (report.get("physics") or {})
            .get("overall_rate"),
        })

    def _key(row: Mapping[str, Any]) -> float:
        v = row["metrics"].get("CVRMSE")
        return float(v) if v is not None else math.inf

    ranking = [r["experiment_id"] for r in sorted(rows, key=_key)]
    return _finish(ctx, make_envelope(
        tool, status="OK",
        inputs={"experiment_ids": [str(e) for e in experiment_ids]},
        summary={
            "experiments": rows,
            "ranking_by_cvrmse": ranking,
            "best": ranking[0] if ranking else None,
            "skipped": skipped,
        },
    ))


def tf_model_publish(
    ctx: ToolContext,
    experiment_id: str,
    *,
    model_id: str,
    version: str,
    description: str | None = None,
    default_out_of_range: str = "reject",
    run_smoke: bool = True,
) -> dict[str, Any]:
    """从已完成实验构建模型包并走发布门禁（model-package §8）。

    成功 → production；任一硬门禁失败 → ok=False + 对应 TFM-10xx 诊断，
    门禁明细可见（entry.last_gate_results）。
    """
    tool = "tf_model_publish"
    inputs = {"experiment_id": experiment_id, "model_id": model_id,
              "version": version}
    try:
        entity = ctx.ledger.get(experiment_id)
        report_path = (ctx.research_root / "experiments" / experiment_id
                       / "report.json")
        if not report_path.exists():
            raise ValueError(f"实验无 report.json（未执行？）: {experiment_id}")
        report = _read_json(report_path)
        if report.get("status") != "completed":
            raise ValueError(f"实验未完成，不得发布: {experiment_id}")
        goal_id = report["goal_id"]
        goal = ctx.ledger.get(goal_id)
        acceptance = (goal.get("definition") or {}).get("acceptance") or {}
        view_entity = ctx.ledger.get(report["dataset_view"])
        view_doc = dict(view_entity["definition"])
        features = [str(f) for f in view_doc["features"]]
        target = str(view_doc["target"])
        dataset_ref = str(view_doc["dataset"])
    except (KeyError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs=inputs)

    # 签名：property_code + 单位来自数据版本（两层命名，DD-11）
    variables = ctx.vault.load_variables(dataset_ref)
    unit_of = {v["property_code"]: v["unit"] for v in variables}

    def _unit(prop: str) -> str:
        if prop not in unit_of:
            raise VaultError("TFV-802", f"签名属性不在数据版本中: {prop}")
        return str(unit_of[prop])

    signature = {
        "model_id": model_id,
        "version": version,
        "object_model": (goal.get("definition") or {}).get("object_model"),
        "inputs": [{"property_code": f, "unit": _unit(f), "dtype": "float",
                    "required": True} for f in features],
        "outputs": [{"property_code": target, "unit": _unit(target),
                     "dtype": "float"}],
    }
    bounds = {
        v["property_code"]: (v.get("min_value"), v.get("max_value"))
        for v in variables
    }
    constraints = {
        "inputs": [{
            "property_code": f,
            "min_value": bounds.get(f, (None, None))[0],
            "max_value": bounds.get(f, (None, None))[1],
            "out_of_range": default_out_of_range,
        } for f in features],
    }
    # golden 输入：物化 View 的工况边界行（§8.2，覆盖边界而非随机抽样）
    mv = materialize_view(ctx.vault, view_doc, ctx.view_cache_root)
    golden_inputs = boundary_rows(mv.table.to_pandas(), features)
    exp_dir = ctx.research_root / "experiments" / experiment_id
    environment = _read_json(exp_dir / "environment.json")
    validation_doc = {
        "physics": report.get("physics"),
        "split": _read_json(exp_dir / "split.json"),
    }
    dataset_lineage = {
        "dataset": dataset_ref,
        "view_id": report["dataset_view"],
        "view_hash": (report.get("conclusion") or {}).get("view_hash"),
    }
    research_lineage = {
        "goal_id": goal_id,
        "hypothesis_id": report.get("hypothesis_id"),
        "experiment_id": experiment_id,
    }
    try:
        ctx.artifacts_root.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(
                dir=ctx.artifacts_root, prefix="pkg_") as staging:
            pkg = build_model_package(
                Path(staging) / "pkg",
                model_id=model_id, version=version, signature=signature,
                artifact_dir=exp_dir / "model",
                metrics=report.get("metrics") or {},
                validation=validation_doc,
                dataset_lineage=dataset_lineage,
                research_lineage=research_lineage,
                environment=environment,
                golden_inputs=golden_inputs,
                constraints=constraints,
                goal_id=goal_id, experiment_id=experiment_id,
                description=description,
            )
            entry = ctx.registry.register(pkg, actor=ctx.actor)
        result = ctx.registry.publish(
            model_id, version, actor=ctx.actor, acceptance=acceptance,
            tfom_registry=ctx.tfom_registry, run_smoke=run_smoke,
        )
    except (ModelRegistryError, VaultError, ValueError) as exc:
        env = _error_envelope(ctx, tool, exc, inputs=inputs)
        gates = ctx.registry.last_gate_results(model_id, version)
        if gates:
            env["summary"]["gates"] = gates
        return env

    ctx.ledger.register_model(
        model_id, experiment_id, actor=ctx.actor,
        metrics=(report.get("metrics") or {}).get("surfaces", {}),
        artifact_path=str(ctx.registry.package_dir(model_id, version)),
        reason="发布门禁全部通过",
    )
    ctx.ledger.create_decision(
        f"publish {model_id}@{version}", "approve",
        actor=ctx.actor,
        rationale="发布门禁全部通过（machine-checked，DD-13）",
        references=[experiment_id],
    )
    return _finish(ctx, make_envelope(
        tool, id=f"{model_id}@{version}", status="PRODUCTION",
        inputs=inputs,
        summary={
            "model_id": model_id, "version": version,
            "status": result["status"],
            "gates": result["gates"],
            "package_dir": str(ctx.registry.package_dir(model_id, version)),
            "golden_samples": len(golden_inputs),
        },
        artifacts=[_artifact(
            ctx.registry.package_dir(model_id, version) / "checksums.json",
            "checksums")],
    ))


# ---------------------------------------------------------------- 预处理工具（I-49）


def tf_preprocess_propose(
    ctx: ToolContext,
    ruleset: Mapping[str, Any],
) -> dict[str, Any]:
    """提交预处理规则集（Schema 校验 + 参数校验 + 版本化存储）。

    Agent 只能提出规则参数；执行的是规则库中预先注册的确定性变换
    （DD-02 方案 C）。新规则一律 status=proposed，需 human 审批。
    """
    tool = "tf_preprocess_propose"
    try:
        rs = RuleSet.model_validate(dict(ruleset))
        for rule in rs.rules:
            validate_rule_params(rule)  # 参数在 propose 时即校验
        # 工具层提交一律记为 proposed + 当前 actor（不信任输入的 status）
        rs = RuleSet.model_validate({
            **rs.model_dump(mode="json"),
            "rules": [
                {**r, "status": "proposed", "proposer": ctx.actor,
                 "approvals": []}
                for r in rs.model_dump(mode="json")["rules"]
            ],
        })
        ctx.preprocess_store.save(rs)
    except (ValidationError, PreprocessError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc,
                               inputs={"ruleset_id": ruleset.get("ruleset_id")})
    return _finish(ctx, make_envelope(
        tool, id=rs.ref, status="PROPOSED",
        inputs={"ruleset_id": rs.ruleset_id, "version": rs.version},
        summary={
            "ruleset": rs.ref,
            "content_hash": rs.content_hash(),
            "rules": [
                {"rule_id": r.rule_id, "rule_type": r.rule_type,
                 "sheet": r.sheet, "status": r.status}
                for r in rs.rules
            ],
        },
    ))


def tf_preprocess_approve(
    ctx: ToolContext,
    ruleset_id: str,
    *,
    version: int | None = None,
    rule_ids: Sequence[str] | None = None,
    note: str | None = None,
) -> dict[str, Any]:
    """审批规则（actor 必须为 human，全部留痕）。"""
    tool = "tf_preprocess_approve"
    inputs = {"ruleset_id": ruleset_id, "version": version,
              "rule_ids": list(rule_ids) if rule_ids else None}
    if ctx.actor != "human":
        return _error_envelope(
            ctx, tool,
            PreprocessError(
                TFPP_APPROVAL_FORBIDDEN,
                f"审批 actor 必须为 human，当前: {ctx.actor!r}",
            ),
            inputs=inputs,
        )
    try:
        rs = ctx.preprocess_store.load(ruleset_id, version)
        targets = set(rule_ids) if rule_ids else {r.rule_id for r in rs.rules}
        unknown = targets - {r.rule_id for r in rs.rules}
        if unknown:
            raise PreprocessError("TFPP-003", f"规则不存在: {sorted(unknown)}")
        now = datetime.now(timezone.utc).isoformat()
        changed = []
        for rule in rs.rules:
            if rule.rule_id not in targets or rule.status == "approved":
                continue
            rule.status = "approved"
            rule.approvals.append(ApprovalRecord(
                actor=ctx.actor, action="approve", at=now, note=note,
            ))
            changed.append(rule.rule_id)
        ctx.preprocess_store.save(rs)  # 内容哈希不变，仅留痕元数据更新
    except (PreprocessError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs=inputs)
    return _finish(ctx, make_envelope(
        tool, id=rs.ref, status="APPROVED" if changed else "UNCHANGED",
        inputs=inputs,
        summary={
            "approved": changed,
            "rules": [
                {"rule_id": r.rule_id, "status": r.status,
                 "approvals": [a.model_dump(mode="json")
                               for a in r.approvals]}
                for r in rs.rules
            ],
        },
    ))


def tf_preprocess_apply(
    ctx: ToolContext,
    path: str | Path,
    ruleset_id: str,
    *,
    version: int | None = None,
    import_into_vault: bool = False,
    write_workbook: bool = False,
    output_path: str | Path | None = None,
) -> dict[str, Any]:
    """执行规则集：原始工作簿 + 规则集 → 内存表模型（可选落 xlsx / 落 vault）。

    审批门禁：`import_into_vault=True`（产生 vault revision 的正式导入）
    要求全部规则 approved；预览执行（不落 vault）允许 proposed 规则。
    逐规则记录 lineage（rule_id、参数、输入/输出指纹）。
    """
    tool = "tf_preprocess_apply"
    inputs = {"path": str(path), "ruleset_id": ruleset_id, "version": version,
              "import_into_vault": import_into_vault}
    try:
        rs = ctx.preprocess_store.load(ruleset_id, version)
        sheets = load_workbook_model(path)
        executions = apply_ruleset(
            sheets, rs, require_approved=import_into_vault,
        )
    except (PreprocessError, ValueError) as exc:
        return _error_envelope(ctx, tool, exc, inputs=inputs)

    run_dir = ctx.research_root / "preprocess" / "runs" / (
        datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%f")
    )
    run_dir.mkdir(parents=True, exist_ok=True)
    lineage_doc = {
        "ruleset": rs.ref,
        "content_hash": rs.content_hash(),
        "source_path": str(path),
        "executed_at": datetime.now(timezone.utc).isoformat(),
        "actor": ctx.actor,
        "rules": [e.to_dict() for e in executions],
    }
    artifacts: list[dict[str, Any]] = []
    lineage_path = run_dir / "lineage.json"
    fd, tmp = tempfile.mkstemp(dir=run_dir, prefix=".tmp_", suffix=".json")
    with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as fp:
        json.dump(lineage_doc, fp, ensure_ascii=False, sort_keys=True, indent=2)
        fp.write("\n")
    os.replace(tmp, lineage_path)
    artifacts.append(_artifact(lineage_path, "preprocess_lineage"))

    dataset_ref = None
    diagnostics: list[dict[str, Any]] = []
    if write_workbook:
        out = Path(output_path) if output_path else (
            run_dir / "processed.xlsx"
        )
        _write_workbook_model(sheets, out)
        artifacts.append(_artifact(out, "processed_workbook"))
    if import_into_vault:
        registry = ctx.tfom_registry or default_registry()
        conv = convert_legacy_tables(
            sheets, registry, source_name=Path(path).name,
        )
        result = import_parsed(
            conv.dataset, conv.timestamps, conv.columns,
            registry=registry,
            pre_diagnostics=conv.pre_diagnostics,
            degradations=conv.degradations,
            source_path=path,
        )
        diagnostics = diagnostic_dicts(result.diagnostics)
        if not result.ok:
            return _finish(ctx, make_envelope(
                tool, ok=False, status="FAILED", inputs=inputs,
                summary={"error": "导入存在 ERROR 级诊断，未写入 vault",
                         "ruleset": rs.ref},
                diagnostics=diagnostics,
                artifacts=artifacts,
            ))
        lineage = conv.lineage(path)
        lineage["preprocess"] = lineage_doc
        try:
            dataset_ref = ctx.vault.store(result, source_path=path,
                                          lineage=lineage)
        except VaultError as exc:
            return _error_envelope(ctx, tool, exc, inputs=inputs)
    return _finish(ctx, make_envelope(
        tool, ok=True, id=dataset_ref or rs.ref,
        status="IMPORTED" if dataset_ref else "APPLIED",
        inputs=inputs,
        summary={
            "ruleset": rs.ref,
            "content_hash": rs.content_hash(),
            "dataset_ref": dataset_ref,
            "rules": [
                {"rule_id": e.rule_id, "rule_type": e.rule_type,
                 "skipped": e.skipped,
                 "input_sha256": (e.input_sha256 or "")[:16],
                 "output_sha256": (e.output_sha256 or "")[:16],
                 "detail": e.detail}
                for e in executions
            ],
        },
        diagnostics=diagnostics,
        artifacts=artifacts,
    ))


def _write_workbook_model(sheets: Mapping[str, Any], path: Path) -> None:
    """内存表模型 → xlsx（openpyxl 写模式；仅 apply 显式要求时调用）。"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, sheet in sheets.items():
        ws = wb.create_sheet(name)
        for row in [*sheet.header_rows, *sheet.data_rows]:
            ws.append(list(row))
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=".tmp_", suffix=".xlsx")
    os.close(fd)
    wb.save(tmp)
    os.replace(tmp, path)


def tf_preprocess_list(ctx: ToolContext) -> dict[str, Any]:
    """列出全部规则集及规则审批状态。"""
    rulesets = ctx.preprocess_store.list_rulesets()
    return _finish(ctx, make_envelope(
        "tf_preprocess_list", status="OK",
        summary={"rulesets": rulesets, "count": len(rulesets)},
    ))


# ---------------------------------------------------------------- 模型实验室（DD-02 方案 C）


def tf_lab_submit(
    ctx: ToolContext,
    name: str,
    source: str,
    *,
    description: str | None = None,
) -> dict[str, Any]:
    """提交模型实验室模块：静态扫描 + 子进程结构校验 + 版本化入库。

    Agent 用它在内置闭集之外起草新模型（协议见 thermoforge_models.lab：
    MODEL_FORMAT / INPUT_ROLES / build_model / load_model）。**通过五连检
    即 status=validated，可立刻在实验里以 category=lab 引用
    （hyperparameters.lab = name 或 name@vN），不需要任何人审批**——
    看指标、改代码、再提交新版本，闭环由 agent 自己走完。
    与最新版内容一致的重交是幂等的（不新增版本）。
    """
    tool = "tf_lab_submit"
    inputs = {"name": name}
    violations = scan_source(source)
    if violations:
        return _finish(ctx, make_envelope(
            tool, ok=False, status="FAILED", inputs=inputs,
            summary={"error": "静态扫描不通过：按 violations 逐条改源码后重交",
                     "violations": violations},
            diagnostics=[{"code": "TFML-001", "level": "ERROR",
                          "count": len(violations),
                          "message": "; ".join(violations)[:500]}],
        ))
    try:
        report: dict[str, Any] | None = None
        with tempfile.TemporaryDirectory(prefix="lab_validate_") as tmp:
            candidate = Path(tmp) / f"{name}.py"
            with open(candidate, "w", encoding="utf-8", newline="\n") as fp:
                fp.write(source)
            report = validate_module(candidate, Path(tmp) / "check")
        if not report.get("ok"):
            raise LabError("TFML-002", "结构校验未通过")
        record = ctx.lab_store.submit(
            name, source, description=description,
            proposer=ctx.actor, validation=report,
        )
    except (LabError, ValueError) as exc:
        env = _error_envelope(ctx, tool, exc, inputs=inputs)
        if report:
            env["summary"]["validation"] = report
        return env
    ref = f"{name}@v{record['version']}"
    src_path = ctx.lab_store.dir / f"{name}.v{record['version']}.py"
    return _finish(ctx, make_envelope(
        tool, id=ref, status="VALIDATED", inputs=inputs,
        summary={
            "ref": ref,
            "content_hash": record["content_hash"],
            "status": record["status"],
            "validation": record.get("validation"),
            "next": f"可直接开实验：model.category=\"lab\" + "
                    f"hyperparameters.lab=\"{ref}\"（无需审批）。"
                    "看完指标要改模型就改源码重交，会自动进下一个版本",
        },
        artifacts=[_artifact(src_path, "lab_source")],
    ))


def tf_lab_deprecate(
    ctx: ToolContext,
    name: str,
    *,
    version: int | None = None,
    note: str | None = None,
) -> dict[str, Any]:
    """停用实验室模块：此后不得再被实验引用（留痕记录停用者与理由）。

    这是实验室通路上唯一的否决口，且不阻塞任何一轮循环——agent 自己也
    该用它清掉已证伪的方案，让规划上下文里只剩还活着的候选。已跑完的
    实验不受影响（源码快照冻结在实验目录里，结论仍可复现）。
    """
    tool = "tf_lab_deprecate"
    inputs = {"name": name, "version": version}
    try:
        record = ctx.lab_store.deprecate(name, version, actor=ctx.actor,
                                         note=note)
    except LabError as exc:
        return _error_envelope(ctx, tool, exc, inputs=inputs)
    return _finish(ctx, make_envelope(
        tool, id=f"{record['name']}@v{record['version']}", status="DEPRECATED",
        inputs=inputs,
        summary={
            "ref": f"{record['name']}@v{record['version']}",
            "content_hash": record["content_hash"],
            "audit": record["audit"],
        },
    ))


def tf_lab_get(
    ctx: ToolContext,
    name: str,
    *,
    version: int | None = None,
) -> dict[str, Any]:
    """读取实验室模块的元数据、校验报告与源码（审批复核用）。"""
    tool = "tf_lab_get"
    inputs = {"name": name, "version": version}
    try:
        record = ctx.lab_store.get(name, version)
    except LabError as exc:
        return _error_envelope(ctx, tool, exc, inputs=inputs)
    src_path = ctx.lab_store.dir / f"{name}.v{record['version']}.py"
    return _finish(ctx, make_envelope(
        tool, id=f"{record['name']}@v{record['version']}", status="OK",
        inputs=inputs,
        summary={
            "ref": f"{record['name']}@v{record['version']}",
            "status": record["status"],
            "proposer": record.get("proposer"),
            "description": record.get("description"),
            "content_hash": record["content_hash"],
            "runnable": record["status"] == "validated",
            "validation": record.get("validation"),
            "audit": record.get("audit") or record.get("approvals") or [],
            "source": record["source"],
        },
        artifacts=[_artifact(src_path, "lab_source")],
    ))


def tf_lab_list(ctx: ToolContext) -> dict[str, Any]:
    """列出全部实验室模块及其校验/停用状态（runnable=true 的可直接引用）。"""
    modules = ctx.lab_store.list()
    return _finish(ctx, make_envelope(
        "tf_lab_list", status="OK",
        summary={
            "modules": modules,
            "count": len(modules),
            "runnable": [m["ref"] for m in modules if m["runnable"]],
        },
    ))


# ---------------------------------------------------------------- 文献工具


def tf_literature_search(
    ctx: ToolContext,
    query: str,
    *,
    sources: Sequence[str] = litsearch.DEFAULT_SOURCES,
    limit: int = 10,
    year_from: int | None = None,
) -> dict[str, Any]:
    """文献调研：跨 CrossRef/arXiv/Semantic Scholar 检索论文（标题/摘要/DOI/引用数）。

    免费 API 无需 key；只读无副作用。单源失败不拖垮整体（status=PARTIAL +
    TFL-001 诊断），全部源不可用才 ok=False（TFL-002）。引用结果时必须带
    DOI/URL，不得编造文献。
    """
    tool = "tf_literature_search"
    limit = max(1, min(int(limit), litsearch.MAX_LIMIT))
    inputs = {"query": query, "sources": list(sources),
              "limit": limit, "year_from": year_from}
    if not (query or "").strip():
        return _finish(ctx, make_envelope(
            tool, ok=False, status="FAILED", inputs=inputs,
            summary={"error": "query 不能为空"},
        ))
    unknown = [s for s in sources if s not in litsearch.SOURCES]
    if unknown:
        return _finish(ctx, make_envelope(
            tool, ok=False, status="FAILED", inputs=inputs,
            summary={"error": f"未知文献源: {unknown}"
                              f"（允许 {list(litsearch.SOURCES)}）"},
            diagnostics=[{"code": litsearch.TFL_SOURCE_UNKNOWN,
                          "level": "ERROR", "count": 1,
                          "message": f"未知文献源: {unknown}"}],
        ))
    try:
        results, failures = litsearch.search(
            query, sources=tuple(sources), limit=limit, year_from=year_from)
    except litsearch.AllSourcesFailed as exc:
        return _finish(ctx, make_envelope(
            tool, ok=False, status="FAILED", inputs=inputs,
            summary={"error": "所有文献源均不可用（网络/限流），"
                              "可稍后重试或换 narrower 查询",
                     "failures": exc.failures},
            diagnostics=[{"code": litsearch.TFL_ALL_SOURCES_FAILED,
                          "level": "ERROR", "count": len(exc.failures),
                          "message": str(exc)[:500]}],
        ))
    failed_sources = {f["source"] for f in failures}
    return _finish(ctx, make_envelope(
        tool, status="PARTIAL" if failures else "OK", inputs=inputs,
        summary={
            "query": query,
            "count": len(results),
            "sources_ok": [s for s in sources if s not in failed_sources],
            "sources_failed": [f["source"] for f in failures],
            "results": results,
        },
        diagnostics=[{"code": litsearch.TFL_SOURCE_FAILED, "level": "WARN",
                      "count": 1, "location": f["source"],
                      "message": f["error"][:500]} for f in failures],
    ))


# architecture §6 工具清单 → 入口（与 harness/tools.json 保持一致）
TOOL_REGISTRY: dict[str, Callable[..., dict[str, Any]]] = {
    "tf_dataset_import": tf_dataset_import,
    "tf_dataset_list": tf_dataset_list,
    "tf_dataset_get": tf_dataset_get,
    "tf_dataset_schema": tf_dataset_schema,
    "tf_dataset_profile": tf_dataset_profile,
    "tf_dataset_query": tf_dataset_query,
    "tf_dataset_sample": tf_dataset_sample,
    "tf_dataset_materialize": tf_dataset_materialize,
    "tf_dataset_compare": tf_dataset_compare,
    "tf_dataset_derive": tf_dataset_derive,
    "tf_dataset_modelability": tf_dataset_modelability,
    "tf_goal_create": tf_goal_create,
    "tf_goal_get": tf_goal_get,
    "tf_research_status": tf_research_status,
    "tf_usage_overview": tf_usage_overview,
    "tf_goal_usage": tf_goal_usage,
    "tf_experiment_usage": tf_experiment_usage,
    "tf_hypothesis_create": tf_hypothesis_create,
    "tf_experiment_plan": tf_experiment_plan,
    "tf_experiment_run": tf_experiment_run,
    "tf_experiment_get": tf_experiment_get,
    "tf_model_compare": tf_model_compare,
    "tf_model_publish": tf_model_publish,
    "tf_preprocess_propose": tf_preprocess_propose,
    "tf_preprocess_approve": tf_preprocess_approve,
    "tf_preprocess_apply": tf_preprocess_apply,
    "tf_preprocess_list": tf_preprocess_list,
    "tf_lab_submit": tf_lab_submit,
    "tf_lab_deprecate": tf_lab_deprecate,
    "tf_lab_get": tf_lab_get,
    "tf_lab_list": tf_lab_list,
    "tf_literature_search": tf_literature_search,
}
