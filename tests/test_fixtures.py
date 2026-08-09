"""TFDC-XLSX fixture 测试。

- 合法 fixture 能被契约模型完整接收，且变量单位与 chiller.v1 物模型一致
  （roadmap Phase 0 验收：示例 Excel 能准确映射到 TFOM）。
- 每个非法变体确实包含其声称的缺陷，期望错误码已登记。
- fixture 目录与 expected_codes.json 保持一致。

注意：Phase 0 尚无导入器，这里验证的是 fixture 本身的构造正确性；
错误码的端到端断言由 Phase 1 导入器测试接管。
"""

import json
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from thermoforge_core.contracts import ObjectModel, TfdcDataset
from thermoforge_core.errors import ERROR_REGISTRY
from thermoforge_core.timeutil import parse_timestamp
from thermoforge_core.units import UnknownUnitError, normalize_unit

import yaml

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"
CONTRACTS_DIR = Path(__file__).resolve().parents[1] / "contracts"


def _expected() -> dict[str, list[str]]:
    with open(FIXTURES_DIR / "expected_codes.json", encoding="utf-8") as fp:
        return json.load(fp)


def _load(name: str):
    return load_workbook(FIXTURES_DIR / name, read_only=False, data_only=True)


def _sheet_rows(wb, sheet: str) -> list[dict]:
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, row)) for row in rows[1:]]


def test_fixture_files_match_expected_codes():
    expected = _expected()
    on_disk = {p.name for p in FIXTURES_DIR.glob("*.xlsx")}
    assert set(expected) == on_disk
    for codes in expected.values():
        for code in codes:
            assert code in ERROR_REGISTRY, f"未登记的错误码: {code}"


def test_valid_fixture_parses_into_contracts():
    wb = _load("valid_minimal.xlsx")
    assert set(wb.sheetnames) == {
        "manifest", "objects", "variables", "data", "parameters", "relations", "bindings",
    }
    manifest = {r["key"]: r["value"] for r in _sheet_rows(wb, "manifest")}
    doc = {
        "manifest": manifest,
        "objects": [
            {k: v for k, v in r.items() if v is not None} for r in _sheet_rows(wb, "objects")
        ],
        "variables": [
            {k: v for k, v in r.items() if v is not None} for r in _sheet_rows(wb, "variables")
        ],
        "parameters": _sheet_rows(wb, "parameters"),
        "relations": [
            {k: v for k, v in r.items() if v is not None} for r in _sheet_rows(wb, "relations")
        ],
        "bindings": _sheet_rows(wb, "bindings"),
    }
    dataset = TfdcDataset.model_validate(doc)
    assert dataset.manifest.dataset_id == "DC01_2026_CHILLER"

    # data 表：timestamp 必须是文本且带时区（conventions.md §3.1）
    ws = wb["data"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    assert header[0] == "timestamp"
    declared = {v.variable_id for v in dataset.variables}
    assert set(header[1:]) == declared
    for row in rows[1:]:
        ts = row[0]
        assert isinstance(ts, str), "timestamp 列必须是文本格式单元格"
        parse_timestamp(ts)  # 带偏移量，可解析为 UTC


def test_valid_fixture_units_match_tfom():
    # 示例 Excel 的变量单位与 chiller.v1 物模型一致
    with open(CONTRACTS_DIR / "tfom" / "examples" / "chiller.v1.yaml", encoding="utf-8") as fp:
        tfom = ObjectModel.model_validate(yaml.safe_load(fp))
    wb = _load("valid_minimal.xlsx")
    for record in _sheet_rows(wb, "variables"):
        prop = tfom.properties[record["property_code"]]
        assert normalize_unit(record["unit"]) == prop.unit


# ---- 非法变体：逐一确认缺陷确实存在 ----

def test_invalid_103_missing_sheet():
    wb = _load("invalid_TFDC-103_missing_sheet.xlsx")
    assert "variables" not in wb.sheetnames


def test_invalid_105_merged_cell():
    wb = _load("invalid_TFDC-105_merged_cell.xlsx")
    assert len(wb["objects"].merged_cells.ranges) > 0


def test_invalid_107_duplicated_header():
    wb = _load("invalid_TFDC-107_duplicated_header.xlsx")
    header = [c.value for c in next(wb["data"].iter_rows(max_row=1))]
    assert len(header) != len(set(header))


def test_invalid_108_formula():
    # 公式检测需 data_only=False 才能看到公式文本（implementation-notes §1.1）
    wb = load_workbook(
        FIXTURES_DIR / "invalid_TFDC-108_formula.xlsx", data_only=False
    )
    formulas = [
        c.value
        for row in wb["data"].iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    assert formulas


def test_invalid_204_bad_timezone():
    wb = _load("invalid_TFDC-204_bad_timezone.xlsx")
    manifest = {r["key"]: r["value"] for r in _sheet_rows(wb, "manifest")}
    assert manifest["timezone"] == "CST"


def test_invalid_205_bad_resolution():
    wb = _load("invalid_TFDC-205_bad_resolution.xlsx")
    manifest = {r["key"]: r["value"] for r in _sheet_rows(wb, "manifest")}
    assert manifest["time_resolution"] == "PT60S"


def test_invalid_304_malformed_variable_id():
    wb = _load("invalid_TFDC-304_malformed_variable_id.xlsx")
    header = [c.value for c in next(wb["data"].iter_rows(max_row=1))]
    assert "CH-01.Evap_Chw_Supply_Temp" in header


def test_invalid_305_undeclared_variable():
    wb = _load("invalid_TFDC-305_undeclared_variable.xlsx")
    data_header = {c.value for c in next(wb["data"].iter_rows(max_row=1))} - {"timestamp"}
    declared = {r["variable_id"] for r in _sheet_rows(wb, "variables")}
    assert data_header - declared == {"CH-01.cw_supply_temp"}


def test_invalid_307_duplicated_variable():
    wb = _load("invalid_TFDC-307_duplicated_variable.xlsx")
    ids = [r["variable_id"] for r in _sheet_rows(wb, "variables")]
    assert len(ids) != len(set(ids))


def test_invalid_401_unknown_unit():
    wb = _load("invalid_TFDC-401_unknown_unit.xlsx")
    units = {r["unit"] for r in _sheet_rows(wb, "variables")}
    with pytest.raises(UnknownUnitError):
        for u in units:
            normalize_unit(u)


def test_invalid_502_naive_timestamp():
    wb = _load("invalid_TFDC-502_naive_timestamp.xlsx")
    ws = wb["data"]
    first_ts = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0][0]
    # Excel 日期类型单元格读出 datetime 而非文本 → TFDC-502
    assert isinstance(first_ts, datetime)


def test_invalid_503_duplicated_timestamp():
    wb = _load("invalid_TFDC-503_duplicated_timestamp.xlsx")
    rows = list(wb["data"].iter_rows(min_row=2, values_only=True))
    timestamps = [r[0] for r in rows]
    assert len(timestamps) != len(set(timestamps))
    # 重复时间戳且值不一致 → ERROR 分支
    dup = next(t for t in timestamps if timestamps.count(t) > 1)
    values = [r[1:] for r in rows if r[0] == dup]
    assert len(set(values)) > 1
