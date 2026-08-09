"""Phase 1 导入器端到端测试：12 份非法 fixture 断言精确错误码/级别/location。

Phase 0 的 tests/test_fixtures.py 只验证缺陷存在；这里验证导入器报出
conventions.md §7 规定的精确诊断。期望码见 tests/fixtures/expected_codes.json。
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from thermoforge_core.errors import Level
from thermoforge_data.importer import ImportOptions, import_xlsx

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"


def _expected() -> dict[str, list[str]]:
    with open(FIXTURES_DIR / "expected_codes.json", encoding="utf-8") as fp:
        return json.load(fp)


@pytest.mark.parametrize(
    "filename,codes",
    sorted((k, v) for k, v in _expected().items() if k.startswith("invalid_")),
    ids=lambda x: x if isinstance(x, str) else ",".join(x),
)
def test_invalid_fixture_reports_exact_codes(filename: str, codes: list[str]):
    result = import_xlsx(FIXTURES_DIR / filename)
    assert not result.ok
    error_codes = [d.code for d in result.diagnostics if d.level == Level.ERROR]
    for code in codes:
        assert code in error_codes, (
            f"{filename} 应报 {code}，实际诊断: "
            f"{[(d.code, d.level.value, d.location) for d in result.diagnostics]}"
        )
    # 诊断必须有结构化 location 与聚合计数
    for d in result.diagnostics:
        assert d.location
        assert d.count >= 1


# ---- 逐变体的精确级别与 location 断言 ----

def _diag(result, code):
    return [d for d in result.diagnostics if d.code == code]


def test_103_missing_sheet():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-103_missing_sheet.xlsx")
    (d,) = _diag(r, "TFDC-103")
    assert d.level == Level.ERROR
    assert d.location == "variables"


def test_105_merged_cell():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-105_merged_cell.xlsx")
    (d,) = _diag(r, "TFDC-105")
    assert d.level == Level.ERROR
    assert d.location == "objects"
    assert d.count >= 1


def test_107_duplicated_header():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-107_duplicated_header.xlsx")
    (d,) = _diag(r, "TFDC-107")
    assert d.level == Level.ERROR
    assert d.location == "data"


def test_108_formula():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-108_formula.xlsx")
    (d,) = _diag(r, "TFDC-108")
    assert d.level == Level.ERROR
    assert d.location == "data"


def test_204_bad_timezone():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-204_bad_timezone.xlsx")
    (d,) = _diag(r, "TFDC-204")
    assert d.level == Level.ERROR
    assert d.location == "manifest"


def test_205_bad_resolution():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-205_bad_resolution.xlsx")
    (d,) = _diag(r, "TFDC-205")
    assert d.level == Level.ERROR
    assert d.location == "manifest"


def test_304_malformed_variable_id():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-304_malformed_variable_id.xlsx")
    (d,) = _diag(r, "TFDC-304")
    assert d.level == Level.ERROR
    assert d.location == "data!CH-01.Evap_Chw_Supply_Temp"


def test_305_undeclared_variable():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-305_undeclared_variable.xlsx")
    (d,) = _diag(r, "TFDC-305")
    assert d.level == Level.ERROR
    assert d.location == "data!CH-01.cw_supply_temp"


def test_307_duplicated_variable():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-307_duplicated_variable.xlsx")
    (d,) = _diag(r, "TFDC-307")
    assert d.level == Level.ERROR
    assert d.location.startswith("variables!")


def test_401_unknown_unit():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-401_unknown_unit.xlsx")
    (d,) = _diag(r, "TFDC-401")
    assert d.level == Level.ERROR
    assert d.location.startswith("variables!")


def test_502_naive_timestamp_error_by_default():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-502_naive_timestamp.xlsx")
    (d,) = _diag(r, "TFDC-502")
    assert d.level == Level.ERROR
    assert d.location == "data!timestamp"
    assert d.count == 4  # 聚合计数（§7.1）


def test_503_duplicated_timestamp_conflict_is_error():
    r = import_xlsx(FIXTURES_DIR / "invalid_TFDC-503_duplicated_timestamp.xlsx")
    (d,) = _diag(r, "TFDC-503")
    assert d.level == Level.ERROR  # 值不一致 → ERROR 分支（§3.5）


# ---- 合法 fixture 与降级路径 ----

def test_valid_fixture_imports_clean():
    r = import_xlsx(FIXTURES_DIR / "valid_minimal.xlsx")
    assert r.ok, [d for d in r.diagnostics]
    assert r.table is not None
    assert r.table.num_rows == 4
    assert str(r.table.schema.field("timestamp").type) == "timestamp[us, tz=UTC]"
    assert r.dataset.manifest.dataset_id == "DC01_2026_CHILLER"


def test_502_naive_timestamp_degraded_with_manifest_tz():
    """显式授权后 naive 时间戳按 manifest.timezone 解释并记录降级（§3.1）。"""
    r = import_xlsx(
        FIXTURES_DIR / "invalid_TFDC-502_naive_timestamp.xlsx",
        options=ImportOptions(allow_naive_with_manifest_tz=True),
    )
    assert r.ok, [d for d in r.diagnostics]
    (d,) = _diag(r, "TFDC-502")
    assert d.level == Level.WARN
    assert d.count == 4
    assert r.degradations  # 降级行为必须记录
    # Asia/Shanghai +08:00：本地 2026-01-01 00:00 → UTC 2025-12-31 16:00
    ts = r.table.column("timestamp").to_pylist()
    assert ts[0].isoformat() == "2025-12-31T16:00:00+00:00"


def test_503_identical_duplicates_dedup_warn(tmp_path):
    """重复时间戳但各列值一致 → 去重保留一条，WARN（§3.5）。"""
    src = load_workbook(FIXTURES_DIR / "valid_minimal.xlsx")
    ws = src["data"]
    row2 = [c.value for c in ws[2]]
    ws.append(row2)
    for cell in ws[ws.max_row]:
        if cell.column == 1 and isinstance(cell.value, str):
            cell.number_format = "@"
    path = tmp_path / "dup_identical.xlsx"
    src.save(path)
    r = import_xlsx(path)
    assert r.ok, [d for d in r.diagnostics]
    (d,) = _diag(r, "TFDC-503")
    assert d.level == Level.WARN
    assert r.table.num_rows == 4  # 去重后仍为 4 行


def test_504_out_of_order_sorted(tmp_path):
    src = load_workbook(FIXTURES_DIR / "valid_minimal.xlsx")
    ws = src["data"]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    rows[1], rows[2] = rows[2], rows[1]  # 交换两行制造乱序
    wb = Workbook()
    wb.remove(wb.active)
    for name in src.sheetnames:
        if name == "data":
            continue
        ws_new = wb.create_sheet(name)
        for row in src[name].iter_rows(values_only=True):
            ws_new.append(list(row))
    ws_data = wb.create_sheet("data")
    ws_data.append([c.value for c in src["data"][1]])
    for row in rows:
        ws_data.append(row)
        for cell in ws_data[ws_data.max_row]:
            if cell.column == 1 and isinstance(cell.value, str):
                cell.number_format = "@"
    path = tmp_path / "out_of_order.xlsx"
    wb.save(path)
    r = import_xlsx(path)
    assert r.ok, [d for d in r.diagnostics]
    (d,) = _diag(r, "TFDC-504")
    assert d.level == Level.WARN
    ts = r.table.column("timestamp").to_pylist()
    assert ts == sorted(ts)  # 已按时间排序
