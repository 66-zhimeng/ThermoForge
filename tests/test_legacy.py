"""旧格式适配器单元测试（小号合成多级表头工作簿）。

真实 43MB 工作簿的集成测试见 tests/test_legacy_integration.py（slow）。
"""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from thermoforge_core.errors import Level
from thermoforge_data.importer import TfomRegistry, default_registry, import_parsed
from thermoforge_data.legacy import convert_legacy_workbook

N_ROWS = 12  # 1 个布尔坏值 / 12 行 = 8.3% < 10% 升级阈值
BASE = datetime(2025, 1, 1, 0, 0, 0)  # naive Excel 序列号（Asia/Shanghai 语义）


def _build_small_legacy(path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("冷水主机")
    ws.append(["设备类", "冷水主机", None, None, None])
    ws.append(["设备名称", "chiller_01", None, None, None])
    ws.append(["物模型", "chiller", None, None, None])
    ws.append(["物模型属性", "power", "电流百分比", "status_run", "load"])
    for i in range(N_ROWS):
        current = 50.0 + i
        ws.append([
            BASE + timedelta(minutes=15 * i),
            300.0 + i,      # power
            current,        # 电流百分比 → current_percent
            1 if i != 5 else 19,  # status_run；第 6 行是坏值 → TFDC-404
            current * 96.72,      # load = current_percent × 96.72（TFDC-604 一致）
        ])

    ws = wb.create_sheet("环境参数")
    ws.append(["设备类", "环境参数", None])
    ws.append(["设备实例", "environment_parameters", None])
    ws.append(["物模型", "environment_parameters", None])
    ws.append(["物模型属性", "ambient_t", "ambient_h"])
    for i in range(N_ROWS):
        ws.append([BASE + timedelta(minutes=15 * i), 25.0 + i * 0.1, 60.0])

    wb.save(path)
    return path


@pytest.fixture()
def small_legacy(tmp_path) -> Path:
    return _build_small_legacy(tmp_path / "legacy_small.xlsx")


def test_convert_small_legacy(small_legacy):
    conv = convert_legacy_workbook(small_legacy, default_registry())
    assert conv.dataset.manifest.timezone == "Asia/Shanghai"
    assert conv.property_mapping == {"电流百分比": "current_percent"}
    variable_ids = {v.variable_id for v in conv.dataset.variables}
    assert "chiller_01.current_percent" in variable_ids
    assert "chiller_01.power" in variable_ids
    assert "environment_parameters.ambient_t" in variable_ids
    assert len(conv.timestamps) == N_ROWS
    # naive 序列号按 Asia/Shanghai 解释：本地 00:00 → UTC 前一日 16:00
    assert conv.timestamps[0].isoformat() == "2024-12-31T16:00:00+00:00"
    # 降级行为记录（TFDC-502 WARN）
    (d,) = conv.pre_diagnostics
    assert d.code == "TFDC-502"
    assert d.level == Level.WARN
    assert conv.degradations
    # lineage 含映射表
    lineage = conv.lineage(small_legacy)
    assert lineage["property_name_mapping"] == {"电流百分比": "current_percent"}


def test_small_legacy_through_pipeline(small_legacy):
    """转换产出走标准导入管线：布尔坏值 REJECT，派生量一致性通过。"""
    conv = convert_legacy_workbook(small_legacy, default_registry())
    result = import_parsed(
        conv.dataset, conv.timestamps, conv.columns,
        pre_diagnostics=conv.pre_diagnostics,
        degradations=conv.degradations,
        source_path=small_legacy,
    )
    assert result.ok, [d for d in result.diagnostics]
    codes = {(d.code, d.level.value, d.location) for d in result.diagnostics}
    # chiller_01.status_run 第 6 行值 19 → TFDC-404 REJECT（聚合 1 次）
    rejects = [d for d in result.diagnostics if d.code == "TFDC-404"]
    assert len(rejects) == 1
    assert rejects[0].level == Level.REJECT
    assert rejects[0].location == "data!chiller_01.status_run"
    assert rejects[0].count == 1
    # TFDC-604：load = current_percent × 96.72 精确成立 → 无派生量诊断
    assert not any(d.code == "TFDC-604" for d in result.diagnostics)
    # 布尔转换：0/1 → False/True
    status = result.table.column("chiller_01.status_run").to_pylist()
    assert status[0] is True
    assert status[5] is None  # 坏值被剔除
    # 幂等：导入结果的 TFDC-502 降级诊断仍在报告中
    assert any(d.code == "TFDC-502" for d in result.diagnostics)
