# -*- coding: utf-8 -*-
"""Inspect workbook structure relevant to filling the 冷却塔 sheet."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

PATH = "data/数据处理_算法导入训练_0723_WX_已加入表冷器.xlsx"

wb = load_workbook(PATH, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)

def dump_headers(ws, ncols=None):
    print(f"\n=== {ws.title} ===")
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=5), start=1):
        vals = [c.value for c in row]
        if ncols:
            vals = vals[:ncols]
        print(f"r{r}:", vals)

for name in ["冷却塔", "冷却水总管", "冷却塔风机"]:
    dump_headers(wb[name])

# check cached values: first data rows of 冷却水总管 and 冷却塔风机
for name in ["冷却水总管", "冷却塔风机"]:
    ws = wb[name]
    print(f"\n--- {name} sample data rows 5-8 ---")
    for row in ws.iter_rows(min_row=5, max_row=8):
        print([c.value for c in row])

wb.close()
